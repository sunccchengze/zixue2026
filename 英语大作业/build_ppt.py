from pathlib import Path
PROJECT = Path(__file__).resolve().parent
# -*- coding: utf-8 -*-
"""Deck engine v2 — editorial style: sharp corners, real gradients, layered images,
oversized type, one dominant accent colour per slide."""
import os, math
import numpy as np
from PIL import Image, ImageDraw, ImageFilter

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

W, H = 13.333, 7.5
BG_W, BG_H = 2200, int(2200 * H / W)
FONT = "Microsoft YaHei"

INK    = (6, 14, 23)
NAVY   = (10, 22, 36)
NAVY2  = (15, 32, 52)
STEEL  = (24, 48, 72)
CYAN   = (56, 198, 224)
TEAL   = (38, 176, 190)
AMBER  = (243, 180, 62)
GREEN  = (54, 200, 128)
RED    = (234, 103, 76)
WHITE  = (255, 255, 255)
PAPER  = (238, 244, 250)
MUTED  = (163, 186, 208)
DIM    = (112, 141, 168)
HAIR   = (72, 104, 134)

IMG  = str(PROJECT / "ppt_images")
FIGH = str(PROJECT / "papers/figures/hires")
TMP  = str(PROJECT / "build/tmp")
os.makedirs(TMP, exist_ok=True)

# ============================================================ imaging
def save_img(im, out, q=90):
    out = str(out)
    if out.lower().endswith((".jpg", ".jpeg")):
        im.convert("RGB").save(out, quality=q, optimize=True, progressive=True)
    else:
        im.save(out, optimize=True)
    return out

def cover_crop(im, ar, anchor=0.42):
    w, h = im.size
    if w / h > ar:
        nw = int(h * ar); x0 = int((w - nw) * 0.5); im = im.crop((x0, 0, x0 + nw, h))
    else:
        nh = int(w / ar); y0 = int((h - nh) * anchor); im = im.crop((0, y0, w, y0 + nh))
    return im

def grad_alpha(shape, direction, a0, a1, soft=1.0):
    h, w = shape
    yy = np.linspace(0.0, 1.0, h)[:, None]
    xx = np.linspace(0.0, 1.0, w)[None, :]
    if direction == "bottom":   t = yy
    elif direction == "top":    t = 1 - yy
    elif direction == "right":  t = xx
    elif direction == "left":   t = 1 - xx
    elif direction == "radial":
        t = np.clip(np.sqrt(((xx - .5) / .78) ** 2 + ((yy - .52) / .78) ** 2), 0, 1)
    else:                       t = np.full((h, w), .5)
    return a0 + (a1 - a0) * np.clip(t, 0, 1) ** soft

def rich_mask(src, out, ar, direction="bottom", a0=0.04, a1=0.92, soft=1.7,
              color=NAVY, base=2000, vignette=0.42, tone=0.16, blur=0.0,
              warm=(255, 216, 158), cool=(34, 104, 168), hi_burn=0.0, spots=None):
    """Split-toned, vignetted photo with a *visible* gradient scrim."""
    im = cover_crop(Image.open(src).convert("RGB"), ar)
    im = im.resize((base, max(1, int(round(base / ar)))), Image.LANCZOS)
    if blur: im = im.filter(ImageFilter.GaussianBlur(blur))
    a = np.asarray(im).astype(np.float32)
    lum = a.mean(2, keepdims=True)
    # split tone: shadows -> cool, highlights -> warm
    tt = (lum / 255.0) ** 1.15
    a = a * (1 - tone) + (np.array(cool, np.float32) * (1 - tt) +
                          np.array(warm, np.float32) * tt) * tone
    # highlight burn towards the scrim colour (keeps white type safe)
    if hi_burn:
        hl = np.clip((lum - 165) / 90.0, 0, 1)
        a = a * (1 - hl * hi_burn) + np.array(color, np.float32)[None, None, :] * hl * hi_burn
    # vignette
    h, w, _ = a.shape
    yy = np.linspace(-1, 1, h)[:, None]; xx = np.linspace(-1, 1, w)[None, :]
    r = np.sqrt((xx * .92) ** 2 + (yy * .92) ** 2)
    a *= (1 - vignette * np.clip(r - 0.30, 0, None) / 0.95)[..., None]
    # local hotspot control: soft pools of shade where type has to stay legible
    if spots:
        yy2 = np.linspace(0, 1, h)[:, None]; xx2 = np.linspace(0, 1, w)[None, :]
        d = np.zeros((h, w), np.float32)
        for cx, cy, rad, st in spots:
            sg = max(rad / 1.7, 1e-3)
            d = np.maximum(d, st * np.exp(-(((xx2 - cx) ** 2 + (yy2 - cy) ** 2) / (2 * sg ** 2))))
        d = d[..., None]
        a = a * (1 - d) + np.array(color, np.float32)[None, None, :] * d
    # gradient scrim
    al = grad_alpha((h, w), direction, a0, a1, soft)[..., None]
    res = a * (1 - al) + np.array(color, np.float32)[None, None, :] * al
    return save_img(Image.fromarray(np.clip(res, 0, 255).astype(np.uint8)), out, q=88)

def scrim_png(out, direction, a0, a1, color=NAVY, soft=1.0):
    w, h = 1500, int(1500 / (W / H))
    a = grad_alpha((h, w), direction, a0, a1, soft)
    px = np.zeros((h, w, 4), np.uint8)
    px[:, :, :3] = np.array(color, np.uint8)
    px[:, :, 3] = (a * 255).astype(np.uint8)
    Image.fromarray(px, "RGBA").save(out)
    return out

def background(out, dots=True, glow=0.0, glow_color=CYAN, base=NAVY, rule=True):
    w, h = BG_W, BG_H
    img = Image.new("RGB", (w, h), base)
    yy = np.linspace(0, 1, h)[:, None]; xx = np.linspace(0, 1, w)[None, :]
    arr = np.asarray(img).astype(np.float32)
    arr = arr + (30 * (0.62 * yy + 0.38 * xx)).astype(np.float32)[..., None]
    if glow:
        rad = np.sqrt(((xx - .14) / .58) ** 2 + ((yy - .06) / .66) ** 2)
        g = np.clip(1 - rad, 0, 1) ** 2 * glow
        arr = arr * (1 - g[..., None]) + np.array(glow_color, np.float32)[None, None, :] * g[..., None]
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
    d = ImageDraw.Draw(img, "RGBA")
    if dots:
        step = int(h / 30); r = max(1, int(h / 950))
        for iy in range(step, h, step):
            for ix in range(step, w, step):
                d.ellipse([ix - r, iy - r, ix + r, iy + r],
                          fill=(128, 172, 210, int(30 * (1 - iy / h * .6))))
    if rule:
        d.line([(0, int(h * .118)), (w, int(h * .118))], fill=(96, 136, 176, 26), width=1)
        d.line([(0, int(h * .905)), (w, int(h * .905))], fill=(96, 136, 176, 22), width=1)
    return save_img(img, out, q=92)

# ============================================================ pptx low level
def _el(t): return OxmlElement(t)

_SPPR = ["xfrm", "prstGeom", "custGeom", "noFill", "solidFill", "gradFill",
         "blipFill", "pattFill", "grpFill", "ln", "effectLst", "effectDag"]

def _sppr_insert(spPr, el, tag):
    idx = len(spPr)
    for i, ch in enumerate(spPr):
        t = ch.tag.split("}")[-1]
        if t in _SPPR and _SPPR.index(t) > _SPPR.index(tag):
            idx = i; break
    spPr.insert(idx, el)

def set_alpha(fill, pct):
    sf = fill._xPr.find(qn("a:solidFill"))
    if sf is None: return
    clr = sf.find(qn("a:srgbClr"))
    if clr is None: return
    for a in clr.findall(qn("a:alpha")): clr.remove(a)
    al = _el("a:alpha"); al.set("val", str(int(max(0, min(100, pct)) * 1000)))
    clr.append(al)

def add_shadow(shape, blur=16, dist=6, alpha=42, color=(0, 0, 0), dir_=3150000):
    spPr = shape._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = _el("a:effectLst"); _sppr_insert(spPr, eff, "effectLst")
    o = _el("a:outerShdw")
    o.set("blurRad", str(int(blur * 12700))); o.set("dist", str(int(dist * 12700)))
    o.set("dir", str(dir_)); o.set("algn", "tl"); o.set("rotWithShape", "0")
    c = _el("a:srgbClr"); c.set("val", "%02X%02X%02X" % color)
    a = _el("a:alpha"); a.set("val", str(int(alpha * 1000))); c.append(a)
    o.append(c); eff.append(o)
    return shape

_RPR = ["ln", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "noFill",
        "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill",
        "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]

def set_font(run, name=FONT):
    rPr = run.font._rPr
    for tag in ("a:latin", "a:ea", "a:cs"):
        old = rPr.find(qn(tag))
        if old is not None: rPr.remove(old)
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = _el(tag); e.set("typeface", name)
        idx = len(rPr)
        for i, ch in enumerate(rPr):
            t = ch.tag.split("}")[-1]
            if "a:" + t in _RPR and _RPR.index("a:" + t) > _RPR.index(tag):
                idx = i; break
        rPr.insert(idx, e)

def tbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf

def para(tf, text, size, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=0, space_before=0, line=None, italic=False,
         spacing=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    if line is not None: p.line_spacing = line
    r = p.add_run(); r.text = " ".join(text) if spacing else text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = RGBColor(*color); set_font(r)
    return p

# ============================================================ shapes
def rect(slide, x, y, w, h, color=NAVY2, alpha=None, shape=MSO_SHAPE.RECTANGLE,
         line_w=None, line_color=None, radius=None, shadow=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if radius is not None and sh.adjustments:
        sh.adjustments[0] = radius
    if color is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*color)
        if alpha is not None: set_alpha(sh.fill, alpha)
    if line_w:
        sh.line.color.rgb = RGBColor(*(line_color or color)); sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    if shadow: add_shadow(sh, **shadow)
    return sh

def grect(slide, x, y, w, h, c1, c2, angle=45, shape=MSO_SHAPE.RECTANGLE,
          line_w=None, line_color=None, radius=None, shadow=None):
    """Native PowerPoint gradient fill — genuinely visible."""
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if radius is not None and sh.adjustments:
        sh.adjustments[0] = radius
    f = sh.fill; f.gradient()
    st = f.gradient_stops
    st[0].color.rgb = RGBColor(*c1); st[0].position = 0.0
    st[-1].color.rgb = RGBColor(*c2); st[-1].position = 1.0
    f.gradient_angle = angle
    if line_w:
        sh.line.color.rgb = RGBColor(*(line_color or c1)); sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    if shadow: add_shadow(sh, **shadow)
    return sh

def hair(slide, x, y, w, color=HAIR, thick=0.9, alpha=None, vertical=False):
    if vertical:
        return rect(slide, x, y, thick / 72.0, w, color=color, alpha=alpha)
    return rect(slide, x, y, w, thick / 72.0, color=color, alpha=alpha)

def pic(slide, path, x, y, w=None, h=None):
    im = Image.open(path); iw, ih = im.size
    if w is None: w = h * iw / ih
    if h is None: h = w * ih / iw
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))

def cover_pic(slide, path, x, y, w, h, anchor=0.42):
    im = Image.open(path); iw, ih = im.size
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    ar, cur = w / h, iw / ih
    if cur > ar:
        k = ar / cur; p.crop_left = p.crop_right = (1 - k) / 2
    else:
        k = cur / ar; p.crop_top = (1 - k) * anchor; p.crop_bottom = (1 - k) - p.crop_top
    return p

def offset_frame(slide, x, y, w, h, color=AMBER, off=0.14, lw=1.1, radius=None):
    """Thin outline offset behind/below an image — cheap, effective depth cue."""
    return rect(slide, x + off, y + off, w, h, color=None,
                shape=(MSO_SHAPE.RECTANGLE if radius is None else MSO_SHAPE.ROUNDED_RECTANGLE),
                line_w=lw, line_color=color, radius=radius)

# ---- plate: white figure card, sharp corners, real shadow ------------------
def plate(slide, path, x, y, w, caption=None, pad=0.13, cap_h=0.0,
          shadow=True, frame_off=None, rot=None):
    im = Image.open(path); iw, ih = im.size
    pw = w - 2 * pad
    ph = pw * ih / iw
    total = pad * 2 + ph + (cap_h if caption else 0.0)
    if frame_off:
        offset_frame(slide, x, y, w, total, color=frame_off[0], off=frame_off[1])
    shp = rect(slide, x, y, w, total, color=PAPER,
               shape=MSO_SHAPE.RECTANGLE,
               line_w=0.9, line_color=(196, 212, 228),
               shadow=dict(blur=18, dist=7, alpha=40) if shadow else None)
    if rot: shp.rotation = rot
    p = pic(slide, path, x + pad, y + pad, w=pw)
    if rot: p.rotation = rot
    if caption:
        tb, tf = tbox(slide, x + pad, y + pad + ph + 0.05, pw, cap_h)
        para(tf, caption, 8.4, (72, 94, 116), first=True, line=1.08)
    return total

# ============================================================ type furniture
def eyebrow(slide, x, y, text, color=CYAN, size=10.5, w=9.0):
    tb, tf = tbox(slide, x, y, w, 0.28)
    para(tf, text, size, color, bold=True, first=True)
    return tb

def headline(slide, x, y, w, text, size=34, color=WHITE, line=1.0, h=None):
    tb, tf = tbox(slide, x, y, w, h or (size / 72.0 * 1.35))
    para(tf, text, size, color, bold=True, first=True, line=line)
    return tb

def rule_bar(slide, x, y, w, color=CYAN, thick=3.0):
    return rect(slide, x, y, w, thick / 72.0, color=color)

def band(slide, y, h, text, accent=CYAN, size=14.5, x=0.9, w=None,
         fill1=None, fill2=None, label=None, big=None, big_color=None):
    """Full-width emphasis band: left accent bar + optional giant number + text."""
    w = w or (W - 1.8)
    grect(slide, x, y, w, h, fill1 or (14, 30, 49), fill2 or (10, 21, 35), angle=0)
    rect(slide, x, y, 0.075, h, color=accent)
    cx = x + 0.32          # matches the 1.22in inner inset used by cards and rows
    if big:
        tb, tf = tbox(slide, cx, y + h / 2 - 0.34, 2.6, 0.68)
        para(tf, big, 30, big_color or accent, bold=True, first=True)
        cx += 2.75
    if label:
        tb, tf = tbox(slide, cx, y + 0.19, 2.4, 0.24)
        para(tf, label, 9.5, accent, bold=True, first=True)
        tb, tf = tbox(slide, cx, y + 0.47, w - (cx - x) - 0.45, h - 0.55)
        para(tf, text, size, (222, 234, 246), first=True, line=1.24)
    else:
        tb, tf = tbox(slide, cx, y, w - (cx - x) - 0.45, h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, text, size, WHITE, bold=(size >= 15), first=True, line=1.22)
    return y + h

def metric_row(slide, x, y, w, h, num, unit, label, note, accent, num_size=40,
               c1=None, c2=None, bar=True, num_w=3.10):
    grect(slide, x, y, w, h, c1 or (14, 31, 51), c2 or (10, 20, 34), angle=0)
    if bar: rect(slide, x, y, 0.07, h, color=accent)
    tb, tf = tbox(slide, x + 0.34, y + h / 2 - 0.42, num_w - 0.20, 0.84)
    para(tf, num, num_size, WHITE, bold=True, first=True, line=0.92)
    p = tf.add_paragraph()
    r = p.add_run(); r.text = unit
    r.font.size = Pt(11.5); r.font.bold = True
    r.font.color.rgb = RGBColor(*accent); set_font(r)
    lx = x + num_w + 0.50; lw = w - num_w - 0.85
    tb, tf = tbox(slide, lx, y + 0.20, lw, 0.34)
    para(tf, label, 12.5, WHITE, bold=True, first=True, line=1.16)
    if note:
        tb, tf = tbox(slide, lx, y + 0.58, lw, h - 0.72)
        para(tf, note, 10.5, MUTED, first=True, line=1.26)

def col_card(slide, x, y, w, h, num, title, body, accent, top_bar=True,
             c1=None, c2=None, num_size=44, title_size=19, body_size=11.5,
             badge=None, badge_y=None, hi=False):
    grect(slide, x, y, w, h, c1 or (13, 29, 48), c2 or (9, 18, 31), angle=90,
          line_w=(1.1 if hi else 0.7), line_color=(accent if hi else (52, 84, 112)))
    if top_bar: rect(slide, x, y, w, 0.065, color=accent)
    tb, tf = tbox(slide, x + 0.28, y + 0.30, w - 0.56, 0.80)
    para(tf, num, num_size, accent, bold=True, first=True, line=0.9)
    tb, tf = tbox(slide, x + 0.28, y + 1.16, w - 0.56, 0.90)
    para(tf, title, title_size, WHITE, bold=True, first=True, line=1.06)
    rule_bar(slide, x + 0.28, y + 2.20, 0.72, accent, 1.8)
    tb, tf = tbox(slide, x + 0.28, y + 2.42, w - 0.56, h - 2.95)
    para(tf, body, body_size, MUTED if not hi else (214, 228, 242), first=True, line=1.30)
    if badge:
        by = badge_y or (y + h - 0.62)
        rect(slide, x + 0.28, by, w - 0.56, 0.36, color=accent)
        tb, tf = tbox(slide, x + 0.28, by, w - 0.56, 0.36, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, badge, 10, (10, 22, 36), bold=True, align=PP_ALIGN.CENTER, first=True)

# ---- charts (native, editable) --------------------------------------------
def vbar(slide, x, y, w, h, cats, series, ymax=100, ystep=20, unit="",
         threshold=None, thr_label=None, val_size=11, barw=0.50):
    lm, bm, tm = 0.60, 0.54, 0.22
    pw, ph = w - lm, h - bm - tm
    for v in range(0, ymax + 1, ystep):
        gy = y + tm + ph * (1 - v / ymax)
        hair(slide, x + lm, gy, pw, (88, 120, 152), 0.7, alpha=55)
        tb, tf = tbox(slide, x + lm - 0.56, gy - 0.11, 0.48, 0.22)
        para(tf, f"{v}", 9, (132, 160, 186), align=PP_ALIGN.RIGHT, first=True)
    rect(slide, x + lm, y + tm + ph, pw, 0.016, color=(140, 172, 200), alpha=170)
    n, ns = len(cats), len(series)
    gw = pw / n
    bw = min(barw, gw * 0.70 / ns)
    for i, c in enumerate(cats):
        for sidx, (nm, col, vals) in enumerate(series):
            v = vals[i]; bh = ph * v / ymax
            bx = x + lm + gw * i + gw / 2 - (bw * ns) / 2 + sidx * bw
            grect(slide, bx, y + tm + ph - bh, bw * 0.86, bh,
                  tuple(min(255, int(c_ * 1.18)) for c_ in col), col, angle=90)
            tb, tf = tbox(slide, bx, y + tm + ph - bh - 0.27, bw * 0.86, 0.24)
            para(tf, f"{v:g}{unit}", val_size, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
        tb, tf = tbox(slide, x + lm + gw * i, y + tm + ph + 0.11, gw, 0.36)
        para(tf, c, 10.5, (200, 218, 234), align=PP_ALIGN.CENTER, first=True, line=1.08)
    if threshold is not None:
        ty = y + tm + ph * (1 - threshold / ymax)
        rect(slide, x + lm, ty, pw, 0.026, color=AMBER, alpha=240)
        tb, tf = tbox(slide, x + lm + pw - 2.70, ty - 0.31, 2.65, 0.28)
        para(tf, thr_label or "", 9.5, AMBER, bold=True, align=PP_ALIGN.RIGHT, first=True)

def full_bar(slide, x, y, w, h, label, value, color, note="", unit="%",
             val_size=30, label_w=2.55):
    """Full-width horizontal bar with a big number parked at the right edge."""
    tb, tf = tbox(slide, x, y + 0.02, label_w, 0.30)
    para(tf, label, 12.5, WHITE, bold=True, align=PP_ALIGN.RIGHT, first=True)
    tx = x + label_w + 0.22
    tw = w - label_w - 1.90
    rect(slide, tx, y + 0.05, tw, h, color=(26, 48, 72), alpha=95)
    bw = tw * value / 100.0
    grect(slide, tx, y + 0.05, max(0.06, bw), h,
          tuple(min(255, int(c * 1.22)) for c in color), color, angle=0)
    tb, tf = tbox(slide, x + w - 1.62, y - 0.09, 1.62, 0.56)
    para(tf, f"{value:g}{unit}", val_size, color, bold=True, align=PP_ALIGN.RIGHT, first=True)
    if note:
        tb, tf = tbox(slide, tx, y + h + 0.07, w - label_w - 0.22, 0.26)
        para(tf, note, 9.5, DIM, first=True)

# ---- page furniture -------------------------------------------------------
def footer(slide, n, total, who, accent=CYAN):
    hair(slide, 0.9, 6.94, W - 1.8, (58, 86, 112), 0.8, alpha=130)
    tb, tf = tbox(slide, 0.9, 7.02, 6.6, 0.24)
    para(tf, "QUALIFICATION REVIEW  ·  WANG ET AL. 2020  ·  ENGINEERING 6, 1115–1121",
         8.4, (104, 133, 160), first=True)
    tb, tf = tbox(slide, W - 4.05, 7.02, 1.55, 0.24)
    para(tf, f"SPEAKER {who}", 9.0, accent, bold=True, align=PP_ALIGN.RIGHT, first=True)
    tb, tf = tbox(slide, W - 2.35, 7.02, 1.45, 0.24)
    para(tf, f"{n:02d} / {total:02d}", 9.5, (150, 178, 204), bold=True,
         align=PP_ALIGN.RIGHT, first=True)

def topline(slide, accent=CYAN):
    rect(slide, 0, 0, W, 0.05, color=accent, alpha=215)

print("engine v2 ready")
