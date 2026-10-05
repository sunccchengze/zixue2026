# -*- coding: utf-8 -*-
"""13-slide deck v2 — varied layouts, sharp corners, visible gradients,
oversized type, layered/bleeding imagery."""
import os
from PIL import Image
from build_ppt import *
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import animations

AR = W / H
M  = "/home/user/build/masks"
F  = "/home/user/build/figs"
os.makedirs(M, exist_ok=True); os.makedirs(F, exist_ok=True)
TOTAL = 13

print("grading imagery ...")
background(f"{TMP}/bg_plain.jpg", dots=True)
background(f"{TMP}/bg_glow.jpg",  dots=True, glow=0.26)

# visible gradients: light end kept near 0 so the photograph reads through
A = {}
A["title"]  = rich_mask(f"{IMG}/01_title_fibers.png",    f"{M}/s01.jpg", AR, "bottom", 0.05, 0.93, 1.55, vignette=.50)
A["crisis"] = rich_mask(f"{IMG}/02_crisis_crowd.png",    f"{M}/s02.jpg", AR, "bottom", 0.06, 0.80, 1.30, vignette=.38)
A["water"]  = rich_mask(f"{IMG}/03_hot_water.png",       f"{M}/s05.jpg", AR, "bottom", 0.05, 0.88, 1.45, vignette=.46,
                        spots=[(0.40, 0.12, 0.62, 0.66), (0.16, 0.30, 0.34, 0.40)])
A["charge"] = rich_mask(f"{IMG}/04_static_charge.png",   f"{M}/s06.jpg", AR, "bottom", 0.06, 0.90, 1.45, vignette=.46)
A["partic"] = rich_mask(f"{IMG}/05_particles_capture.png", f"{M}/s08.jpg", AR, "bottom", 0.06, 0.91, 1.45, vignette=.46)
A["cyc"]    = rich_mask(f"{IMG}/01_title_fibers.png",    f"{M}/s10.jpg", AR, "radial", 0.30, 0.97, 0.85, vignette=.55, blur=1.4)
A["fact"]   = rich_mask(f"{IMG}/06_factory.png",         f"{M}/s11.jpg", 5.70 / 5.45, "bottom", 0.10, 0.90, 1.30, base=1400, vignette=.40)
A["turb"]   = rich_mask(f"{IMG}/07_turbine_blades.png",  f"{M}/s13.jpg", AR, "left",  0.04, 0.94, 1.50, vignette=.50)

S = {}
S["b"] = scrim_png(f"{M}/scrim_b.png", "bottom", 0.00, 0.42)
S["l"] = scrim_png(f"{M}/scrim_l.png", "left",   0.00, 0.46)

# paper figures sized to their real on-slide width (230 dpi)
for nm, inches in [("fig1", 6.32), ("fig2", 4.78), ("fig3", 2.49), ("fig4", 5.44), ("fig5", 1.62)]:
    im = Image.open(f"{FIGH}/{nm}.png").convert("RGB")
    tw = int(inches * 230)
    save_img(im.resize((tw, max(1, int(round(tw * im.size[1] / im.size[0])))), Image.LANCZOS),
             f"{F}/{nm}.jpg", q=92)
print("imagery ready")

prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
def new(): return prs.slides.add_slide(BLANK)

# ================================================== 01 POSTER
s = new()
cover_pic(s, A["title"], 0, 0, W, H); cover_pic(s, S["b"], 0, 0, W, H)
cover_pic(s, S["l"], 0, 0, W, H)
topline(s, CYAN)
tb, tf = tbox(s, 10.92, 1.92, 1.28, 2.10)
para(tf, "01", 104, (30, 68, 96), bold=True, align=PP_ALIGN.RIGHT, first=True, line=0.80)
eyebrow(s, 0.9, 2.02, "R E S E A R C H   P R E S E N T A T I O N   ·   2 0 2 6   A C 0 4", CYAN, 10.5, 10.0)
headline(s, 0.9, 2.42, 10.2, "Can a Disposable Mask\nBe Reused?", 54, line=0.98, h=1.58)
rule_bar(s, 0.9, 4.40, 2.05, AMBER, 4)
tb, tf = tbox(s, 0.9, 4.66, 9.6, 0.80)
para(tf, "A qualification review of hot-water decontamination", 17, (224, 236, 246), first=True, line=1.16)
para(tf, "and charge regeneration", 17, (224, 236, 246), line=1.16)
tb, tf = tbox(s, 0.9, 5.86, 9.6, 0.26)
para(tf, "Wang D. et al.  ·  Engineering 6 (2020) 1115–1121  ·  Beijing University of Chemical Technology",
     11, (158, 186, 212), first=True)
tb, tf = tbox(s, 0.9, 6.28, 9.0, 0.30)
para(tf, "PRESENTED BY   甲   ·   乙   ·   丙   ·   丁", 13.5, WHITE, bold=True, first=True)
hair(s, 12.34, 2.02, 4.30, CYAN, 1.1, vertical=True, alpha=150)
rect(s, 12.17, 2.02, 0.11, 0.11, color=CYAN, alpha=200)
rect(s, 12.17, 6.21, 0.11, 0.11, color=CYAN, alpha=200)

# ================================================== 02 WHY  (metric rows + visible photo)
s = new()
cover_pic(s, A["crisis"], 0, 0, W, H)
topline(s, AMBER)
eyebrow(s, 0.9, 0.48, "0 1   ·   C O N T E X T", AMBER)
headline(s, 0.9, 0.74, 9.6, "Why This Matters", 38)
tb, tf = tbox(s, 0.9, 1.50, 8.8, 0.32)
para(tf, "In early 2020 a disposable mask stopped being disposable — it became a strategic resource.",
     13, MUTED, first=True)
metric_row(s, 0.9, 2.30, 7.60, 1.12, "8", "MILLION MASKS / DAY",
           "Total mask production in China", "Reported for 23 January 2020", AMBER, num_size=42,
           c1=(17, 36, 58), c2=(10, 21, 35))
metric_row(s, 0.9, 3.60, 7.60, 1.12, "1.4", "BILLION PEOPLE",
           "The population that supply had to cover", "The demand–supply gap was enormous", AMBER,
           num_size=42, c1=(17, 36, 58), c2=(10, 21, 35))
metric_row(s, 0.9, 4.90, 7.60, 1.12, "3", "DAYS TO ACCEPT",
           "From submission to acceptance", "Received 26 May · Accepted 29 May 2020", AMBER,
           num_size=42, c1=(17, 36, 58), c2=(10, 21, 35))
tb2 = tbox(s, 12.30, 2.30, 0.32, 4.0)[0]
tb2.rotation = -90
tf2 = tb2.text_frame
para(tf2, "CHINA  ·  JANUARY 2020", 9.5, (168, 198, 224), first=True, align=PP_ALIGN.RIGHT)
band(s, 6.30, 0.58, "So the question became urgent: is there a way to make one mask last longer — safely?",
     AMBER, 15)
footer(s, 2, TOTAL, "甲", AMBER)

# ================================================== 03 THREE QUESTIONS (tall columns)
s = new()
cover_pic(s, f"{TMP}/bg_plain.jpg", 0, 0, W, H)
topline(s, CYAN)
eyebrow(s, 0.9, 0.48, "0 2   ·   F R A M E W O R K", CYAN)
headline(s, 0.9, 0.74, 11.0, "Three Questions for a Qualification Review", 32)
tb, tf = tbox(s, 0.9, 1.52, 11.2, 0.30)
para(tf, "This paper is not a discovery. It is a qualification campaign — and any review must answer three questions.",
     13, MUTED, first=True)
cw, gap = 3.64, 0.30
for i, (n, t, b, c) in enumerate([("01", "KILL?", "Does the procedure actually inactivate the virus?", CYAN),
                                  ("02", "FILTER?", "Does the mask still capture particles afterwards?", AMBER),
                                  ("03", "SURVIVE?", "Does it hold up after real hours of wear?", GREEN)]):
    col_card(s, 0.9 + i * (cw + gap), 2.28, cw, 4.42, n, t, b, c,
             num_size=48, title_size=26, body_size=12)
footer(s, 3, TOTAL, "甲", CYAN)

# ================================================== 04 PROCEDURE (split, figure bleeds right)
s = new()
cover_pic(s, f"{TMP}/bg_plain.jpg", 0, 0, W, H)
topline(s, CYAN)
eyebrow(s, 0.9, 0.48, "0 3   ·   M E T H O D", CYAN)
headline(s, 0.9, 0.74, 6.4, "The Procedure", 34)
tb, tf = tbox(s, 0.9, 1.48, 5.30, 0.66)
para(tf, "Two steps, no chemicals, no special equipment.", 13, MUTED, first=True, line=1.22)
para(tf, "56 °C water, then a hair dryer.", 13, MUTED, line=1.22)
hair(s, 6.46, 1.95, 4.85, HAIR, 1.0, vertical=True, alpha=110)
metric_row(s, 0.9, 2.42, 5.20, 1.55, "56 °C", "FOR 30 MINUTES", "STEP 01  ·  SOAK",
           "The envelope protein of the virus denatures and the virus loses its ability to infect.",
           CYAN, num_size=31, num_w=2.35, c1=(15, 33, 54), c2=(9, 19, 32))
metric_row(s, 0.9, 4.22, 5.20, 1.55, "10", "MINUTES, HAIR DRYER", "STEP 02  ·  DRY",
           "Recharges the melt-blown filter layer with static electricity.",
           AMBER, num_size=31, num_w=2.35, c1=(15, 33, 54), c2=(9, 19, 32))
band(s, 5.98, 0.66, "NO CHEMICALS   ·   NO SPECIAL EQUIPMENT   ·   DOABLE IN A KITCHEN",
     CYAN, 13, x=0.9, w=5.20)
plate(s, f"{F}/fig1.jpg", 6.75, 1.90, 6.40,
      "Fig. 1 — Household containers, water temperature over 30 min, and a used mask under 365 nm UV light, "
      "before and after treatment.   Wang et al. (2020).", cap_h=0.34, shadow=True)
rect(s, 6.75, 1.90, 0.06, 4.06, color=CYAN, alpha=200)
footer(s, 4, TOTAL, "乙", CYAN)

# ================================================== 05 TEMPERATURE
s = new()
cover_pic(s, A["water"], 0, 0, W, H)
topline(s, AMBER)
eyebrow(s, 0.9, 0.48, "0 4   ·   M E T H O D", AMBER)
headline(s, 0.9, 0.74, 11.0, "Does the Water Stay Hot Enough?", 31)
grect(s, 0.9, 2.05, 7.30, 4.02, (14, 31, 52), (8, 17, 29), angle=135,
      line_w=0.9, line_color=(58, 90, 120), shadow=dict(blur=20, dist=8, alpha=46))
vbar(s, 1.20, 2.42, 6.55, 3.30,
     cats=["Aluminium\nbasin", "Plastic\nlunch box", "Stainless\nthermos cup"],
     series=[("Freshly boiled", (62, 108, 154), [90, 90, 90]),
             ("After 30 min", CYAN, [60, 60, 85])],
     ymax=100, ystep=20, unit="°C", threshold=56, thr_label="56 °C required", barw=0.46)
tb, tf = tbox(s, 1.20, 5.62, 6.8, 0.26)
para(tf, "WATER TEMPERATURE (°C)", 9.5, (150, 178, 204), first=True)
eyebrow(s, 8.55, 2.05, "THE TARGET", AMBER, 10.5, 3.9)
tb, tf = tbox(s, 8.55, 2.26, 3.90, 1.16)
para(tf, "56", 84, AMBER, bold=True, first=True, line=0.82)
tb, tf = tbox(s, 8.55, 3.46, 3.90, 0.30)
para(tf, "°C — FOR 30 MINUTES", 12.5, WHITE, bold=True, first=True)
rule_bar(s, 8.55, 3.88, 0.80, AMBER, 2.4)
tb, tf = tbox(s, 8.55, 4.14, 3.90, 1.90)
para(tf, "Boiling water starts near 90 °C in all three containers.", 12, (212, 228, 242), first=True,
     line=1.34, space_after=9)
para(tf, "After 30 minutes the basin and the lunch box are still around 60 °C; the thermos holds 85 °C.",
     11.5, MUTED, line=1.34, space_after=9)
para(tf, "The requirement is met with things people already own.", 11.5, MUTED, line=1.34)
band(s, 6.28, 0.58, "Three everyday containers — and the water is still hot enough after half an hour.",
     AMBER, 14)
footer(s, 5, TOTAL, "乙", AMBER)

# ================================================== 06 STATIC CHARGE (full-width bars)
s = new()
cover_pic(s, A["charge"], 0, 0, W, H)
topline(s, CYAN)
eyebrow(s, 0.9, 0.48, "0 5   ·   M E T H O D", CYAN)
headline(s, 0.9, 0.74, 11.6, "The Hidden Problem: Static Charge", 31)
tb, tf = tbox(s, 0.9, 1.46, 11.2, 0.30)
para(tf, "Hot water kills the virus — and it also destroys the charge that makes the mask work.", 13, MUTED, first=True)
full_bar(s, 0.9, 2.42, 7.90, 0.54, "New mask", 100, (86, 122, 160), "", val_size=30)
full_bar(s, 0.9, 3.72, 7.90, 0.54, "Air-dried 10 h", 60, AMBER,
         "too slow; bacteria can grow in a wet mask", val_size=30)
full_bar(s, 0.9, 5.02, 7.90, 0.54, "Hair dryer 10 min", 90, GREEN,
         "fast, and recovers most of the charge", val_size=30)
tb, tf = tbox(s, 0.9, 6.18, 7.9, 0.26)
para(tf, "RECOVERY OF ELECTROSTATIC CHARGE ON THE FILTER LAYER", 9.5, (150, 178, 204), first=True)
eyebrow(s, 9.20, 2.42, "WHY IT MATTERS", CYAN, 10.5, 3.2)
rule_bar(s, 9.20, 2.76, 0.80, CYAN, 2.4)
tb, tf = tbox(s, 9.20, 3.00, 3.30, 3.4)
para(tf, "Electrostatic adsorption is the only mechanism that captures nano-sized particles while keeping "
         "the mask easy to breathe through.", 12, WHITE, first=True, line=1.34, space_after=10)
para(tf, "Hot water removes that charge. So the charge has to be put back — that is what the hair-dryer "
         "step is for.", 11.5, MUTED, line=1.34, space_after=10)
para(tf, "Drying fast also stops bacteria from growing inside a wet mask.", 11.5, MUTED, line=1.34)
band(s, 6.42, 0.52, "The recipe: hot water, 56 °C, 30 minutes; hair dryer, 10 minutes.", CYAN, 13.5)
footer(s, 6, TOTAL, "乙", CYAN)

# ================================================== 07 ANATOMY (layered/bleeding figures)
s = new()
cover_pic(s, f"{TMP}/bg_glow.jpg", 0, 0, W, H)
topline(s, AMBER)
eyebrow(s, 0.9, 0.48, "0 6   ·   S T R U C T U R E", AMBER)
headline(s, 0.9, 0.74, 7.0, "What’s Inside a Mask", 34)
tb, tf = tbox(s, 0.9, 1.48, 5.4, 0.46)
para(tf, "Three layers of non-woven fabric — one of them does almost all of the work.", 13, MUTED, first=True, line=1.22)
def layer_row(y, h, tag, body, accent, hi=False, big=None):
    grect(s, 0.9, y, 5.35, h, (16, 35, 56) if hi else (12, 26, 43),
          (9, 18, 30) if hi else (8, 16, 27), angle=0,
          line_w=(1.1 if hi else 0.7), line_color=(accent if hi else (48, 78, 106)))
    rect(s, 0.9, y, 0.075, h, color=accent)
    tb, tf = tbox(s, 1.22, y + 0.19, 3.0, 0.28)
    para(tf, tag, 11.5, accent, bold=True, first=True)
    if big:
        tb, tf = tbox(s, 1.22, y + 0.52, 4.85, 0.42)
        para(tf, big, 22, WHITE, bold=True, first=True)
        tb, tf = tbox(s, 1.22, y + 1.00, 4.85, h - 1.10)
        para(tf, body, 11.5, (222, 232, 242), first=True, line=1.30)
    else:
        tb, tf = tbox(s, 1.22, y + 0.52, 4.85, h - 0.62)
        para(tf, body, 11.5, MUTED, first=True, line=1.30)
layer_row(2.22, 1.02, "OUTER LAYER", "Waterproof non-woven. Blocks liquid droplets sprayed by others.",
          (76, 116, 158))
layer_row(3.42, 1.48, "FILTER LAYER", "Polypropylene melt-blown, 100–1000 µm thick. This is the working layer.",
          AMBER, hi=True, big="fibres 1–10 µm")
layer_row(5.08, 1.02, "INNER LAYER", "Ordinary non-woven. Absorbs moisture released by the wearer.",
          (76, 116, 158))
band(s, 6.30, 0.58, "The filter layer is where everything happens.", AMBER, 15, x=0.9, w=5.35)
# figures: big one bleeds off the right edge, small one overlaps it
p2 = pic(s, f"{F}/fig2.jpg", 8.55, 0.82, w=4.78)
add_shadow(p2, blur=22, dist=10, alpha=52)
plate(s, f"{F}/fig3.jpg", 6.48, 4.12, 2.75, shadow=True, rot=-2.4, frame_off=(AMBER, 0.12))
tb, tf = tbox(s, 6.48, 6.72, 2.9, 0.24)
para(tf, "Fig. 3 — waterproof test, SEM", 8.2, (120, 148, 176), first=True)
footer(s, 7, TOTAL, "丙", AMBER)

# ================================================== 08 MECHANISMS (five columns)
s = new()
cover_pic(s, A["partic"], 0, 0, W, H)
topline(s, AMBER)
eyebrow(s, 0.9, 0.48, "0 7   ·   M E C H A N I S M", AMBER)
headline(s, 0.9, 0.74, 11.0, "Five Filtration Mechanisms", 33)
tb, tf = tbox(s, 0.9, 1.46, 11.2, 0.30)
para(tf, "A mask does not work like a sieve. Five mechanisms act together — one of them dominates.", 13, MUTED, first=True)
mechs = [("01", "Brownian\ndiffusion", "Random motion drives the smallest particles into fibres."),
         ("02", "Entrapment", "Particles too large to follow the air stream get stuck."),
         ("03", "Inertial\ncollision", "Heavier particles cannot turn with the air and hit a fibre."),
         ("04", "Gravity\nsedimentation", "Large particles simply settle out of the airflow."),
         ("05", "Electrostatic\nadsorption", "Charge pulls nano-sized particles in — at almost no cost to breathing resistance.")]
cw, gap = 2.10, 0.2575
for i, (n, t, b) in enumerate(mechs):
    x = 0.9 + i * (cw + gap); key = (i == 4)
    col_card(s, x, 2.16, cw, 3.70, "", t, b, AMBER, top_bar=True,
             num_size=13, title_size=13, body_size=9.6, hi=key,
             c1=((30, 44, 34) if key else (13, 29, 48)),
             c2=((16, 26, 26) if key else (9, 18, 31)),
             badge=("THE KEY" if key else None), badge_y=(2.16 + 3.70 - 0.56))
    ov = rect(s, x + 0.28, 2.40, 0.50, 0.50, color=(AMBER if key else CYAN),
              shape=MSO_SHAPE.OVAL, alpha=(240 if key else 200))
    tb, tf = tbox(s, x + 0.28, 2.40, 0.50, 0.50, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, n, 12, (12, 26, 42) if key else (8, 18, 30), bold=True, align=PP_ALIGN.CENTER, first=True)
    rule_bar(s, x + 0.24, 3.52, 0.62, AMBER if key else (78, 112, 144), 1.6)
band(s, 6.00, 0.80, "THE ONE THAT MATTERS — Electrostatic adsorption is the only mechanism that captures "
     "nano-sized particles while keeping the mask breathable, and it is exactly the one that hot water destroys.",
     AMBER, 14)
footer(s, 8, TOTAL, "丙", AMBER)

# ================================================== 09 RESULTS
s = new()
cover_pic(s, f"{TMP}/bg_plain.jpg", 0, 0, W, H)
topline(s, GREEN)
eyebrow(s, 0.9, 0.48, "0 8   ·   E V I D E N C E", GREEN)
headline(s, 0.9, 0.74, 8.0, "Test Results vs Standards", 33)
tb, tf = tbox(s, 0.9, 1.46, 6.4, 0.46)
para(tf, "Measured with a TSI 8130 against NaCl aerosol, following the Chinese national standards.", 13, MUTED, first=True, line=1.22)
metric_row(s, 0.9, 2.12, 5.55, 1.24, ">95", "% MEASURED", "Bacterial filtration — BFE",
           "Required ≥ 95%   ·   YY/T 0969–2013 · YY 0469–2011", GREEN, num_size=34, num_w=2.00,
           c1=(15, 33, 54), c2=(9, 19, 32))
metric_row(s, 0.9, 3.50, 5.55, 1.24, "92.3", "% MEASURED", "Particle filtration — surgical",
           "Required ≥ 30%   ·   YY 0469–2011", AMBER, num_size=34, num_w=2.00,
           c1=(15, 33, 54), c2=(9, 19, 32))
metric_row(s, 0.9, 4.88, 5.55, 1.24, ">95", "% MEASURED", "Particle filtration — KN95",
           "Required ≥ 95%   ·   GB 2626–2019", GREEN, num_size=34, num_w=2.00,
           c1=(15, 33, 54), c2=(9, 19, 32))
grect(s, 0.9, 6.26, 5.55, 0.50, (13, 28, 46), (9, 18, 31), angle=0)
tb, tf = tbox(s, 1.10, 6.26, 5.15, 0.50, anchor=MSO_ANCHOR.MIDDLE)
para(tf, "NaCl aerosol · 0.075 µm · 30 L/min (medical, surgical) · 85 L/min (KN95)",
     9.2, (146, 174, 200), first=True)
plate(s, f"{F}/fig4.jpg", 7.63, 1.86, 5.52,
      "Fig. 4 — BFE and PFE of regenerated masks, performance after 10 cycles, and after "
      "sterilisation at 121 °C.   Wang et al. (2020).", cap_h=0.34, shadow=True)
rect(s, 7.63, 1.86, 0.06, 5.02, color=GREEN, alpha=200)
footer(s, 9, TOTAL, "丙", GREEN)

# ================================================== 10 TEN CYCLES
s = new()
cover_pic(s, A["cyc"], 0, 0, W, H)
topline(s, CYAN)
eyebrow(s, 0.9, 0.48, "0 9   ·   D U R A B I L I T Y", CYAN)
headline(s, 0.9, 0.74, 11.0, "After Ten Cycles", 34)
tb, tf = tbox(s, -0.45, 1.30, 5.0, 4.6)
para(tf, "10", 300, (30, 76, 108), bold=True, align=PP_ALIGN.CENTER, first=True, line=0.78)
grect(s, 4.55, 2.10, 7.88, 2.16, (14, 31, 52), (8, 17, 29), angle=120,
      line_w=0.9, line_color=(58, 90, 120), shadow=dict(blur=20, dist=8, alpha=46))
rect(s, 4.55, 2.10, 0.075, 2.16, color=CYAN)
tb, tf = tbox(s, 4.95, 2.34, 3.0, 0.30)
para(tf, "TEN FULL SOAK-AND-DRY CYCLES", 10.5, CYAN, bold=True, first=True)
tb, tf = tbox(s, 4.95, 2.74, 7.05, 1.30)
para(tf, "Five brands of masks were put through ten full cycles of hot-water soaking and hair-dryer charging.",
     13, (222, 234, 246), first=True, line=1.34, space_after=8)
para(tf, "Filtration performance barely changed.", 16, CYAN, bold=True, line=1.34, space_after=8)
para(tf, "Ten cycles — essentially a fatigue-life test for a disposable mask.", 12.5, (180, 204, 226),
     italic=True, line=1.30)
grect(s, 4.55, 4.46, 7.88, 2.16, (14, 31, 52), (8, 17, 29), angle=120,
      line_w=0.9, line_color=(58, 90, 120), shadow=dict(blur=20, dist=8, alpha=46))
rect(s, 4.55, 4.46, 0.075, 2.16, color=AMBER)
tb, tf = tbox(s, 4.95, 4.70, 7.0, 0.28)
para(tf, "PUSHED FURTHER — PRESSURISED STEAM AT 121 °C, 30 MIN", 10.5, AMBER, bold=True, first=True)
full_bar(s, 4.95, 5.12, 6.30, 0.40, "3M 1860", 99.2, GREEN, "", val_size=24, label_w=1.30)
full_bar(s, 4.95, 5.74, 6.30, 0.40, "KF94", 96.6, CYAN, "", val_size=24, label_w=1.30)
tb, tf = tbox(s, 11.55, 5.30, 0.80, 0.90)
para(tf, "95%\nmin.", 9.5, AMBER, bold=True, align=PP_ALIGN.RIGHT, first=True, line=1.1)
footer(s, 10, TOTAL, "丙", CYAN)

# ================================================== 11 REAL WEAR (bleeding photo column)
s = new()
cover_pic(s, f"{TMP}/bg_plain.jpg", 0, 0, W, H)
topline(s, AMBER)
eyebrow(s, 0.9, 0.48, "1 0   ·   F I E L D", AMBER)
headline(s, 0.9, 0.74, 11.0, "Real Wear, Real Factory", 33)
tb, tf = tbox(s, 0.9, 1.46, 11.2, 0.30)
para(tf, "Laboratory tests are clean. Real life is not — so the authors ran both.", 13, MUTED, first=True)
pf = pic(s, A["fact"], 0.0, 2.00, w=5.70, h=5.50)
add_shadow(pf, blur=24, dist=0, alpha=0)
rect(s, 0.0, 2.00, 5.70, 0.05, color=AMBER, alpha=220)
tb, tf = tbox(s, 0.55, 4.30, 5.0, 0.86)
para(tf, "122,500", 50, WHITE, bold=True, first=True, line=0.92)
rule_bar(s, 0.55, 5.28, 1.70, AMBER, 3.0)
tb, tf = tbox(s, 0.55, 5.48, 4.85, 1.30)
para(tf, "Zhejiang Runtu Co. · over 4,000 staff", 12.5, WHITE, bold=True, first=True, line=1.30)
para(tf, "20 February – 30 March 2020", 11.5, (212, 226, 240), line=1.30)
para(tf, "One mask per person per day  →  one mask every three days", 11.5, AMBER, bold=True, line=1.30)
rows = [("WATERPROOF", "100 mL in 20 s, vacuum 3 min",
         "After 10 cycles: no seepage. SEM shows the fibres intact.", CYAN),
        ("8 HOURS OF REAL WEAR", "Surgical vs KN95",
         "Surgical fell by 0.5–12% (15 samples). All 10 KN95 stayed above 95%.", AMBER),
        ("THE REAL CULPRIT", "It is not the hot water",
         "Dirt and oil from the skin change the wettability of the fibres.", GREEN)]
yy = 2.02
for i, (tag, sub, body, c) in enumerate(rows):
    rect(s, 6.05, yy, 0.06, 0.90, color=c)
    tb, tf = tbox(s, 6.37, yy + 0.02, 3.98, 0.26)
    para(tf, tag, 11, c, bold=True, first=True)
    tb, tf = tbox(s, 6.37, yy + 0.28, 3.98, 0.26)
    para(tf, sub, 12.5, WHITE, bold=True, first=True)
    tb, tf = tbox(s, 6.37, yy + 0.56, 3.98, 0.36)
    para(tf, body, 10.5, MUTED, first=True, line=1.26)
    if i < 2: hair(s, 6.05, yy + 1.04, 4.30, HAIR, 0.8, alpha=90)
    yy += 1.10
plate(s, f"{F}/fig5.jpg", 10.55, 2.02, 1.88, shadow=True, rot=2.6, frame_off=(AMBER, 0.11))
tb, tf = tbox(s, 10.55, 5.20, 2.0, 0.30)
para(tf, "Fig. 5 — fluorescent\npenetrant inspection", 8.2, (120, 148, 176), first=True, line=1.12)
band(s, 5.82, 1.04, "Deployed at scale: one company cut mask use to a third in five weeks.",
     AMBER, 13.5, x=6.05, w=6.38, label="FROM THE LAB TO THE FACTORY")
footer(s, 11, TOTAL, "丁", AMBER)

# ================================================== 12 VERDICT
s = new()
cover_pic(s, f"{TMP}/bg_plain.jpg", 0, 0, W, H)
topline(s, GREEN)
eyebrow(s, 0.9, 0.48, "1 1   ·   V E R D I C T", GREEN)
headline(s, 0.9, 0.74, 11.0, "Verdict: Approved — With Limits", 33)
tb, tf = tbox(s, 0.9, 1.46, 11.2, 0.30)
para(tf, "What the evidence supports — and what it simply does not.", 13, MUTED, first=True)
hair(s, 6.66, 2.02, 3.30, HAIR, 1.0, vertical=True, alpha=130)
tb, tf = tbox(s, 0.9, 2.05, 5.40, 0.42)
para(tf, "P R O V E N", 25, GREEN, bold=True, first=True)
rule_bar(s, 0.9, 2.60, 1.10, GREEN, 2.6)
tb, tf = tbox(s, 0.9, 2.80, 5.40, 1.40)
para(tf, "Hot water plus a hair dryer keeps the filtration performance of disposable, surgical and KN95 "
         "masks essentially intact for up to 10 cycles — and it was deployed at scale, with measurable savings.",
     12.5, WHITE, first=True, line=1.34)
tb, tf = tbox(s, 6.96, 2.05, 5.47, 0.42)
para(tf, "N O T   P R O V E N", 25, RED, bold=True, first=True)
rule_bar(s, 6.96, 2.60, 1.10, RED, 2.6)
tb, tf = tbox(s, 6.96, 2.80, 5.47, 1.40)
para(tf, "• Virus killing was never measured directly\n"
         "• NaCl particles, not the actual virus\n"
         "• Only six products were tested\n"
         "• Fit and seal were never tested at all",
     11.5, (244, 220, 214), first=True, line=1.42)
band(s, 4.52, 1.02, "In §3.5 the worn KN95 masks were treated with hot water; in the Conclusions the same result "
     "is described as pressurised steam at 121 °C. Two different treatments — the conclusion stands, but "
     "published papers still need careful reading.",
     AMBER, 11.5, label="ONE INCONSISTENCY")
band(s, 5.72, 1.06, "For surgical masks the required PFE is only 30%. So “92.3%” sounds impressive partly "
     "because the bar is low. A good reviewer always asks: impressive compared to what?",
     CYAN, 11.5, label="COMPARED TO WHAT?", big="30%")
footer(s, 12, TOTAL, "丁", GREEN)

# ================================================== 13 CLOSING
s = new()
cover_pic(s, A["turb"], 0, 0, W, H)
topline(s, AMBER)
tb, tf = tbox(s, 9.60, 0.45, 3.4, 2.6)
para(tf, "”", 240, (34, 74, 100), bold=True, align=PP_ALIGN.RIGHT, first=True, line=0.78)
eyebrow(s, 0.9, 1.42, "1 2   ·   W H Y   T H I S   P A P E R", AMBER)
headline(s, 0.9, 1.92, 10.2, "Why We Chose This Paper", 42)
rule_bar(s, 0.9, 2.92, 2.05, AMBER, 4)
tb, tf = tbox(s, 0.9, 3.26, 9.60, 2.62)
para(tf, "This is not a discovery. It is a qualification campaign:", 17, (216, 230, 244), first=True,
     line=1.30, space_after=4)
para(tf, "set the acceptance criteria, run to N cycles, inspect, then release.", 17, WHITE, bold=True,
     line=1.30, space_after=14)
para(tf, "That is exactly what engineers do to turbine blades and filters.", 15, (196, 214, 232),
     line=1.30, space_after=16)
para(tf, "And in an emergency, the best engineering is often not the most advanced —", 16, (232, 240, 248),
     bold=True, line=1.30)
para(tf, "it is the kind that works with a pot of hot water and a hair dryer.", 24, AMBER, bold=True, line=1.26)
tb, tf = tbox(s, 0.9, 6.18, 6.0, 0.32)
para(tf, "T H A N K   Y O U", 15, WHITE, bold=True, first=True)
hair(s, 12.34, 1.42, 4.60, AMBER, 1.1, vertical=True, alpha=150)
footer(s, 13, TOTAL, "丁", AMBER)

out = "/home/user/Mask-Reuse-Qualification-Review.pptx"
n_anim = animations.animate_presentation(prs)
prs.save(out)
print(f"animated {n_anim} objects")
print("saved:", out)
