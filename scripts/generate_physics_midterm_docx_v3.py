#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the full, question-by-question V3 college-physics homework guide."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
COURSE = ROOT / "大学物理" / "课程作业"
OUT = ROOT / "大学物理" / "期中复习" / "大学物理期中_12-15次作业图文精析_孙承泽_2253710052.docx"

sys.path.insert(0, str(COURSE))
from content12 import S12  # noqa: E402
from content13 import S13  # noqa: E402
from content14 import S14  # noqa: E402
from content15 import S15  # noqa: E402
from 逐题深度解析_v3 import ANSWERS, DEEP_SOLUTIONS, ERRATA, QUESTION_OVERRIDES  # noqa: E402
from 计算题解析基线_v3 import CALC_SOLUTIONS  # noqa: E402

SESSIONS = [
    (12, S12, "机械波", "波动方程 · 干涉 · 驻波 · 多普勒"),
    (13, S13, "波动光学 1：干涉", "相干性 · 双缝 · 薄膜 · 劈尖 · 牛顿环 · 迈克尔逊"),
    (14, S14, "波动光学 2：衍射与光栅", "单缝衍射 · 光栅缺级 · 光谱重叠 · 分辨本领"),
    (15, S15, "波动光学 3：偏振", "偏振态 · 马吕斯定律 · 布儒斯特定律 · 双折射与波片"),
]

NAVY = "183B56"
BLUE = "2A6F97"
TEAL = "2A9D8F"
INK = "263238"
GRAY = "5C6770"
LIGHT_BLUE = "EAF3F8"
LIGHT_YELLOW = "FFF4D6"
LIGHT_GREEN = "EAF6F1"
LIGHT_GOLD = "FFF8E8"
WHITE = "FFFFFF"

FIGURES = {
    (12, "choice", 3): ("fig12-wave-propagation.png", "波形沿传播方向平移：可据此判断固定质点的瞬时运动方向。"),
    (12, "choice", 7): ("fig12-two-source-interference.png", "两相干源的波峰/波谷交点：峰峰、谷谷加强，峰谷相消。"),
    (12, "choice", 9): ("fig12-standing-wave.png", "驻波包络：波节不动，波腹振幅最大；相邻波节段反相。"),
    (12, "choice", 10): ("fig12-doppler.png", "声源与观察者相向：接收频率升高。"),
    (12, "blank", 15): ("fig12-two-source-interference.png", "等幅相干波的加强点与减弱点。"),
    (13, "choice", 4): ("fig13-double-slit.png", "杨氏双缝：条纹间距由双缝中心距决定。"),
    (13, "choice", 8): ("fig13-newton-rings.png", "牛顿环为等厚干涉；固定级次对应固定膜厚。"),
    (13, "choice", 9): ("fig13-newton-rings.png", "从下方观察是透射光；透射与反射条纹互补。"),
    (13, "blank", 16): ("fig13-newton-rings.png", "浸液后同级暗环半径按 1/√n 缩小。"),
    (13, "blank", 19): ("fig13-newton-rings.png", "牛顿环接触点的明暗要逐界面数反射相变。"),
    (13, "calc", 22): ("fig13-wedge.png", "劈尖缺陷：等厚条纹弯曲方向指示工件凹凸。"),
    (13, "calc", 24): ("fig13-wedge.png", "圆锥与平板形成轴对称空气劈尖，等厚线为同心圆。"),
    (14, "choice", 1): ("fig14-single-slit.png", "单缝宽度控制中央明纹宽度；主光轴控制中央位置。"),
    (14, "choice", 5): ("fig14-grating.png", "光栅主极大受单缝衍射包络调制，落在单缝暗纹处即缺级。"),
    (14, "choice", 7): ("fig14-resolution.png", "瑞利判据：孔径越大、波长越短，最小分辨角越小。"),
    (14, "blank", 14): ("fig14-single-slit.png", "单缝中央明纹由左右第一级暗纹界定。"),
    (14, "blank", 17): ("fig14-grating.png", "光栅明纹与单缝暗纹重合时发生缺级。"),
    (15, "choice", 1): ("fig15-malus.png", "自然光过偏振片后强度减半；第二片按马吕斯定律继续透射。"),
    (15, "choice", 3): ("fig15-polarization-states.png", "自然光、部分偏振、线偏振和圆偏振的状态示意。"),
    (15, "choice", 4): ("fig15-polarization-states.png", "转动检偏器：圆偏振强度恒定，线偏振可消光。"),
    (15, "blank", 14): ("fig15-malus.png", "自然光先减半，再用 I=I₀cos²θ 求两透光轴夹角。"),
    (15, "blank", 20): ("fig15-waveplate.png", "半波片使两主轴分量相位差改变，可反转椭圆偏振旋向。"),
    (15, "choice", 7): ("fig15-brewster.png", "布儒斯特角下反射光为 s 偏振，电矢量垂直入射面。"),
    (15, "choice", 8): ("fig15-birefringence.png", "单轴晶体 o 光球面波阵面、e 光旋转椭球面波阵面。"),
    (15, "blank", 12): ("fig15-brewster.png", "p 偏振方向平行入射面，在布儒斯特角无反射。"),
    (15, "blank", 19): ("fig15-birefringence.png", "波片利用 o/e 两主轴的相位延迟改变偏振态。"),
}

ORIGINAL_THUMBNAILS = {
    12: ("第十二次-p7.png", "第十二次作业原卷首页（含波动方程与读图题）。"),
    13: ("第十三次-p13.png", "第十三次作业原卷首页（干涉）。"),
    14: ("第十四次-p19.png", "第十四次作业原卷首页（衍射与光栅）。"),
    15: ("第十五次-p25.png", "第十五次作业原卷首页（偏振）。"),
}


def set_run_font(run, size=10.5, bold=False, color=INK, cn="宋体", en="Cambria Math"):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:cs"), en)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for edge, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7E1E8", size="5"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_keep_together(row):
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def add_paragraph(doc, text="", size=10.5, bold=False, color=INK, before=0, after=4,
                  align=None, indent=0.0, font_cn="宋体", keep=False, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.16
    pf.left_indent = Inches(indent) if indent else None
    pf.widow_control = True
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    if text:
        r = p.add_run(text)
        set_run_font(r, size, bold, color, font_cn)
        r.italic = italic
    return p


def add_heading(doc, text, level=1):
    if level == 1:
        p = add_paragraph(doc, text, size=18, bold=True, color=NAVY, before=14, after=8,
                          font_cn="黑体", keep=True)
        ppr = p._p.get_or_add_pPr()
        pbd = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "5")
        bottom.set(qn("w:color"), BLUE)
        pbd.append(bottom)
        ppr.append(pbd)
    elif level == 2:
        add_paragraph(doc, text, size=14, bold=True, color=BLUE, before=10, after=6,
                      font_cn="黑体", keep=True)
    else:
        add_paragraph(doc, text, size=11.2, bold=True, color=NAVY, before=8, after=3,
                      font_cn="黑体", keep=True)


def add_label_line(doc, label, text, bg=None, label_color=BLUE, text_color=INK, size=10.1):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.45)
    cell = table.cell(0, 0)
    set_cell_margins(cell)
    set_cell_borders_color = BLUE if bg == LIGHT_BLUE else ("E1C56F" if bg == LIGHT_YELLOW else "D7E1E8")
    set_table_borders(table, color=set_cell_borders_color, size="4")
    if bg:
        shade_cell(cell, bg)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(label)
    set_run_font(r, size, True, label_color, "黑体")
    r2 = p.add_run(text)
    set_run_font(r2, size, False, text_color, "宋体")
    set_keep_together(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_answer_box(doc, answer):
    add_label_line(doc, "参考答案　", answer, bg=LIGHT_YELLOW,
                   label_color="8A5A00", text_color=INK, size=10.3)


def add_figure(doc, filename, caption, width=5.8):
    path = COURSE / "图卡" / filename
    if not path.exists():
        raise FileNotFoundError(path)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(width)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=50, start=60, bottom=40, end=60)
    set_table_borders(table, color="E0E7EB", size="4")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width - 0.18))
    set_keep_together(table.rows[0])
    add_paragraph(doc, caption, size=8.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER,
                  before=0, after=5, italic=True)


def add_original_page(doc, session):
    filename, caption = ORIGINAL_THUMBNAILS[session]
    path = ROOT / "大学物理" / "期中复习" / "原卷缩略图" / filename
    if not path.exists():
        raise FileNotFoundError(path)
    add_paragraph(doc, "原卷图（首页）", size=10, bold=True, color=BLUE,
                  before=5, after=3, font_cn="黑体", keep=True)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.1)
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
    set_table_borders(table, color="D7E1E8", size="4")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(5.8))
    set_keep_together(table.rows[0])
    add_paragraph(doc, caption, size=8.4, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER,
                  after=5, italic=True)


def clean_text(text):
    return re.sub(r"\*\*(.*?)\*\*", r"\1", str(text)).replace("**", "")


def stem_text(session, category, number, q):
    override = QUESTION_OVERRIDES.get((session, category, number))
    if override:
        return override
    raw = q["stem"].strip()
    prefix = re.match(r"^\s*\d+[.、]\s*", raw)
    return raw[prefix.end():] if prefix else raw


def add_solution_block(doc, session, category, number, q, answer, solution):
    type_names = {"choice": "选择题", "blank": "填空题", "calc": "计算题"}
    add_heading(doc, f"第 {number} 题｜{type_names[category]}", level=3)
    add_label_line(doc, "题目　", stem_text(session, category, number, q), bg=LIGHT_BLUE,
                   label_color=BLUE, size=10.0)
    add_answer_box(doc, answer)

    if category == "calc":
        add_paragraph(doc, "解题思路与分步推导", size=10.2, bold=True, color=BLUE,
                      before=1, after=3, font_cn="黑体", keep=True)
        analysis = clean_text(solution["analysis"])
        lines = [line.strip() for line in analysis.splitlines() if line.strip()]
        for i, line in enumerate(lines, 1):
            add_paragraph(doc, line, size=9.85, color=INK, before=1, after=3,
                          indent=0.12)
        if solution.get("tip"):
            add_label_line(doc, "易错提醒　", clean_text(solution["tip"]), bg=LIGHT_GOLD,
                           label_color="986B00", size=9.6)
    else:
        add_label_line(doc, "解题思路　", solution["idea"], bg=LIGHT_GREEN,
                       label_color=TEAL, size=9.95)
        add_paragraph(doc, "推理步骤", size=10, bold=True, color=BLUE,
                      before=1, after=2, font_cn="黑体", keep=True)
        for i, step in enumerate(solution["steps"], 1):
            add_paragraph(doc, f"步骤 {i}　{step}", size=9.85, color=INK,
                          before=0, after=2, indent=0.12)
        if solution.get("check"):
            add_label_line(doc, "结果自检　", solution["check"], bg=LIGHT_GREEN,
                           label_color=TEAL, size=9.4)
        if solution.get("pitfall"):
            add_label_line(doc, "易错提醒　", solution["pitfall"], bg=LIGHT_GOLD,
                           label_color="986B00", size=9.4)

    figure = FIGURES.get((session, category, number))
    if figure:
        add_figure(doc, *figure, width=5.45)
    # Mild divider keeps consecutive questions visually distinct without forcing page breaks.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    ppr = p._p.get_or_add_pPr()
    pbd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "3")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "E4E9ED")
    pbd.append(bottom)
    ppr.append(pbd)


def add_answer_key(doc, session, data):
    add_heading(doc, "本次答案速查", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(0.58), Inches(2.65), Inches(0.58), Inches(2.65)]
    for i, width in enumerate(widths):
        table.columns[i].width = width
    headers = ["題號", "答案", "題號", "答案"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade_cell(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, 9.2, True, WHITE, "黑體")
    all_answers = []
    for category in ("choice", "blank"):
        for i, answer in enumerate(ANSWERS[session][category], 1):
            num = i if category == "choice" else i + 10
            all_answers.append((num, answer))
    calc_answers = [(n, CALC_SOLUTIONS[(session, n)]["answer"]) for n in range(21, 25)]
    all_answers.extend(calc_answers)
    for i in range(0, len(all_answers), 2):
        row = table.add_row()
        set_keep_together(row)
        for col in range(2):
            if i + col >= len(all_answers):
                continue
            num, answer = all_answers[i + col]
            cnum = row.cells[col * 2]
            cans = row.cells[col * 2 + 1]
            cnum.width = widths[0]
            cans.width = widths[1]
            cnum.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cnum.paragraphs[0].add_run(str(num))
            set_run_font(r, 8.8, True, NAVY, "宋体")
            r2 = cans.paragraphs[0].add_run(clean_text(answer))
            set_run_font(r2, 8.2, False, INK, "宋体")
            for cell in (cnum, cans):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=50, start=70, bottom=50, end=70)
            if i // 2 % 2 == 1:
                shade_cell(cnum, "F6F8FA")
                shade_cell(cans, "F6F8FA")
    add_paragraph(doc, "提示：选择题答案字母以逐题推导与本版勘误为准；遇到旧答案册标号冲突，请看对应题目的推理步骤。",
                  size=8.5, color=GRAY, before=3, after=5, italic=True)


def add_cover(doc):
    for _ in range(3):
        add_paragraph(doc, "", after=10)
    add_paragraph(doc, "大学物理（下册）", size=26, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=8, font_cn="黑体")
    add_paragraph(doc, "期中考试范围 · 全题逐题深度解析", size=21, bold=True, color=BLUE,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=10, font_cn="黑体")
    add_paragraph(doc, "第十二~十五次作业｜V3.0 全题重修版", size=15, bold=True, color=TEAL,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=8, font_cn="黑体")
    add_paragraph(doc, "机械波　·　干涉　·　衍射与光栅　·　偏振", size=12.5, color=GRAY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=28)

    info = [
        ("课程范围", "第十二至第十五次作业，共 96 题"),
        ("逐题结构", "题干要点 → 参考答案 → 解题思路 → 推理步骤 → 自检与易错提醒"),
        ("题型覆盖", "40 道选择题 + 40 道填空题 + 16 道计算题"),
        ("姓名 / 学号", "孙承泽　/　2253710052"),
        ("班级", "能动强基2501"),
        ("修订日期", "2026 年 9 月 28 日"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(5.2)
    for idx, (key, value) in enumerate(info):
        c0, c1 = table.rows[idx].cells
        shade_cell(c0, LIGHT_BLUE)
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = c0.paragraphs[0].add_run(key)
        set_run_font(r0, 10, True, NAVY, "黑体")
        r1 = c1.paragraphs[0].add_run(value)
        set_run_font(r1, 10, False, INK, "宋体")
        for cell in (c0, c1):
            set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_paragraph(doc, "本版重点：不再把多个题号合并成一段；选择题与填空题也逐题说明为什么、怎么做、如何检查。",
                  size=10.5, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER,
                  before=24, after=8)
    doc.add_page_break()


def add_front_matter(doc):
    add_heading(doc, "阅读说明与修订范围", level=1)
    notes = [
        "本册把第十二至第十五次作业的 96 题全部拆成独立题块：40 道选择、40 道填空、16 道计算。每题都给题意、答案、思路和分步推理；选择/填空另加自检与易错提醒。",
        "计算题保留原来已经写清的分步推导；本轮把容易被答案表一句话带过的选择题和填空题全部补成逐题解析，不再以‘答案速查表’代替讲解。",
        "图像题保留原卷首页缩略图，并在关键题旁加上波形、干涉、衍射、偏振示意图。图卡用于解释物理关系；题面图示以原卷为准。",
        "本版对照原卷与答案资料复核了易混标号、数值和单位；有争议处把物理推导写在题目下面，并在本册勘误表列明。",
    ]
    for note in notes:
        add_label_line(doc, "说明　", note, bg=LIGHT_BLUE, label_color=BLUE, size=9.6)
    add_heading(doc, "本册导航", level=2)
    for text in [
        "一　第十二次：机械波（选择 1–10，填空 11–20，计算 21–24）",
        "二　第十三次：波动光学 1·干涉（选择 1–10，填空 11–20，计算 21–24）",
        "三　第十四次：波动光学 2·衍射与光栅（选择 1–10，填空 11–20，计算 21–24）",
        "四　第十五次：波动光学 3·偏振（选择 1–10，填空 11–20，计算 21–24）",
        "附录　答案资料差异与勘误",
    ]:
        add_paragraph(doc, "• " + text, size=10.4, color=INK, after=3, indent=0.08)
    add_heading(doc, "先记住的解题主线", level=2)
    for title, body in [
        ("机械波", "读 ω、k、φ，再判传播方向；图像题用特殊点和微移法；能量题分清质点速度、波速、能流密度。"),
        ("干涉", "先找参与干涉的两束光，逐界面数反射相变，再写光程差与明暗条件。"),
        ("衍射", "单缝暗纹定包络，光栅方程定主极大，二者重合判断缺级；分辨题先算角再换长度。"),
        ("偏振", "自然光先减半；线偏振逐片用马吕斯；布儒斯特题标清 p/s；波片题看相位延迟。"),
    ]:
        add_label_line(doc, f"{title}　", body, bg=LIGHT_GREEN, label_color=TEAL, size=9.8)
    doc.add_page_break()


def build_document():
    # Guard against accidentally skipping any item or reverting to a bundled answer block.
    expected = {(s, c, n) for s in (12, 13, 14, 15)
                for c, start, end in (("choice", 1, 10), ("blank", 11, 20))
                for n in range(start, end + 1)}
    assert set(DEEP_SOLUTIONS) == expected, (
        f"non-calculation solution coverage mismatch: missing={expected-set(DEEP_SOLUTIONS)}, "
        f"extra={set(DEEP_SOLUTIONS)-expected}")
    assert len(CALC_SOLUTIONS) == 16
    assert all((s, n) in CALC_SOLUTIONS for s in (12, 13, 14, 15) for n in range(21, 25))
    assert all(len(ANSWERS[s]["choice"]) == 10 and len(ANSWERS[s]["blank"]) == 10
               for s in (12, 13, 14, 15))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.74)
    sec.right_margin = Inches(0.74)
    sec.header_distance = Inches(0.30)
    sec.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Cambria Math"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15

    # Header and page number.
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    rr = hp.add_run("大学物理 · 第十二至十五次作业全题解析 V3.0")
    set_run_font(rr, 8, False, GRAY, "宋体")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("孙承泽　·　能动强基2501　　|　　第 ")
    set_run_font(r, 8, False, GRAY, "宋体")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(begin)
    r._r.append(instr)
    r._r.append(separate)
    r._r.append(text)
    r._r.append(end)

    add_cover(doc)
    add_front_matter(doc)

    for section_index, (session, data, title, topic_line) in enumerate(SESSIONS):
        if section_index:
            doc.add_page_break()
        add_heading(doc, f"第 {session} 次　{title}", level=1)
        add_paragraph(doc, "本次考点：" + topic_line, size=10.1, color=GRAY,
                      align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
        add_original_page(doc, session)
        add_answer_key(doc, session, data)

        add_heading(doc, "知识框架（先认判据，再逐题做）", level=2)
        for item in data["framework"]:
            head, bullets = item
            add_heading(doc, head, level=3)
            for bullet in bullets:
                add_paragraph(doc, "• " + bullet, size=9.65, color=INK,
                              before=0, after=2, indent=0.08)

        add_heading(doc, "一、选择题 1–10｜每题单独拆解", level=2)
        for i, q in enumerate(data["choice"], 1):
            key = (session, "choice", i)
            add_solution_block(doc, session, "choice", i, q, ANSWERS[session]["choice"][i - 1],
                               DEEP_SOLUTIONS[key])

        add_heading(doc, "二、填空题 11–20｜每空逐项推导", level=2)
        for i, q in enumerate(data["blank"], 11):
            key = (session, "blank", i)
            add_solution_block(doc, session, "blank", i, q, ANSWERS[session]["blank"][i - 11],
                               DEEP_SOLUTIONS[key])

        add_heading(doc, "三、计算题 21–24｜保留完整分步解", level=2)
        for i, q in enumerate(data["calc"], 21):
            calc = CALC_SOLUTIONS[(session, i)]
            # The assignment source carries the complete problem statement; the calculation baseline
            # carries the previously checked answer and step-by-step solution.
            add_solution_block(doc, session, "calc", i, q, calc["answer"], calc)

    doc.add_page_break()
    add_heading(doc, "附录　答案资料差异与勘误", level=1)
    add_paragraph(doc, "下列条目是会直接影响选项或计算结果的易错处。本册正文已经按推导结果统一，列在这里便于对照旧版答案资料。",
                  size=10.2, color=INK, after=7)
    for idx, note in enumerate(ERRATA, 1):
        add_label_line(doc, f"{idx:02d}　", note, bg=LIGHT_GOLD,
                       label_color="986B00", size=9.5)
    add_heading(doc, "资料依据", level=2)
    for ref in [
        "题干：大学物理/课程作业/content12.py 至 content15.py，并对照各次原卷首页缩略图核查图示题。",
        "答案核对：大学物理/资料原件/大学物理下册作业解析.pdf 与逐题物理推导；发现的标号或答案差异已在本册标注。",
        "概念图：大学物理/课程作业/图卡/ 下的机械波、干涉、衍射、偏振示意图；图示用于讲解，不替代原题图。",
    ]:
        add_paragraph(doc, "• " + ref, size=9.3, color=GRAY, after=3)
    add_paragraph(doc, "姓名：孙承泽　　学号：2253710052　　班级：能动强基2501",
                  size=10, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER,
                  before=14, after=6, font_cn="黑体")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Created: {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")
    print("Coverage: 4×(10 choice + 10 blank + 4 calculation) = 96 independently numbered problems")


if __name__ == "__main__":
    build_document()
