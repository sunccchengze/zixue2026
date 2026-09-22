#!/usr/bin/env python3
"""独立复核1-2，不导入生成器、不改写既有图或文档。

运行：.venv/bin/python scripts/verify_mechanics_homework_20260922.py
随机参数仅用于验证受力模型，不是教材题设，也不是官方答案。
几何对应关系仍须人工对照原题；本程序不声称能自动判读图像。
"""
from pathlib import Path
import math
import random
import zipfile

import pymupdf
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '工程力学/作业'


def equilibrium(loads):
    """每项为(x,y,Fx,Fy)；检查合力和关于三个不同矩心的力矩。"""
    scale_f = max(1, sum(abs(fx) + abs(fy) for _, _, fx, fy in loads))
    assert abs(sum(fx for _, _, fx, _ in loads)) < 1e-10 * scale_f
    assert abs(sum(fy for _, _, _, fy in loads)) < 1e-10 * scale_f
    for ox, oy in [(0, 0), (2.7, -1.3), (-5.2, 4.1)]:
        moments = [(x - ox) * fy - (y - oy) * fx for x, y, fx, fy in loads]
        assert abs(sum(moments)) < 1e-10 * max(1, sum(map(abs, moments)))


def verify_models():
    rng = random.Random(20260922)
    for _ in range(120):
        alpha, beta, theta = [rng.uniform(.1, 1.3) for _ in range(3)]
        P, F = rng.uniform(1, 1000), rng.uniform(1, 1000)
        sa, ca = math.sin(alpha), math.cos(alpha)
        # (a) 球心原点，半径1；绳作用点(-sinα, cosα)，墙接触点(-1,0)。
        T = P / ca
        equilibrium([(0, 0, 0, -P), (-sa, ca, -T*sa, T*ca), (-1, 0, P*math.tan(alpha), 0)])
        # (b) A与B两条径向法线；右侧尖角的法线并非水平或竖直。
        cb, sb = math.cos(beta), math.sin(beta)
        NA, NB = P*cb/math.cos(alpha-beta), P*sa/math.cos(alpha-beta)
        assert NA > 0 and NB > 0
        equilibrium([(0, 0, 0, -P), (-sa, -ca, NA*sa, NA*ca), (cb, -sb, -NB*cb, NB*sb)])
        # (c) P2=(-H,-V)，保留其向左、向下的两个分量，作用在同一点。
        L = rng.uniform(2, 10)
        x1, x2 = sorted([rng.uniform(.1*L, .9*L) for _ in range(2)])
        H, V = rng.uniform(1, 100), rng.uniform(1, 100)
        Ny = (P*x1 + V*x2) / L
        Nx = -Ny*math.tan(alpha)
        equilibrium([(0, 0, H-Nx, P+V-Ny), (x1, 0, 0, -P), (x2, 0, -H, -V), (L, 0, Nx, Ny)])
        # (d) 定滑轮的两项切线张力+轴反力；动滑轮两张力+挂物连接力。
        ct, st = math.cos(theta), math.sin(theta)
        equilibrium([(-st, ct, -F*ct, -F*st), (1, 0, 0, -F), (0, 0, F*ct, F*(1+st))])
        equilibrium([(-1, 0, 0, F), (1, 0, 0, F), (0, 0, 0, -2*F)])
        # (e) AD=d，AB=b，BC=c；P作用点距B为ell；q仅作用AD。
        d, b, c = 2., 3.5, rng.uniform(1, 5)
        ell, q = rng.uniform(.1*c, .9*c), rng.uniform(1, 100)
        Cy = P*ell/c
        Cx = -Cy*math.tan(alpha)
        Bx, By = -Cx, P-Cy  # BC上的内铰力，AB上必须取反。
        ND = (q*d*d/2 + By*b)/d
        Ay, Ax = q*d+By-ND, Bx
        ab = [(0, 0, Ax, Ay), (d, 0, 0, ND), (d/2, 0, 0, -q*d), (b, 0, -Bx, -By)]
        bc = [(b, 0, Bx, By), (b+ell, 0, 0, -P), (b+c, 0, Cx, Cy)]
        equilibrium(ab)
        equilibrium(bc)
        equilibrium(ab[:-1] + bc[1:])  # 整体删除内铰力仍平衡。
    print('PASS：120组参数×7个分离体及整体，合力/三个矩心力矩均通过；未复用生成器验算函数')


def verify_artifacts():
    book = pymupdf.open(ROOT / '工程力学/框架版教材/工程力学（框架汇编版）.pdf')
    original = pymupdf.open(ROOT / '工程力学/资料原件/工程力学（第三版）_1-130.pdf')
    for i, j in [(29, 32), (30, 33)]:
        assert book[i].get_text() == original[j].get_text()
        assert book[i].get_pixmap(dpi=120).samples == original[j].get_pixmap(dpi=120).samples
    path = OUT / '2026-09-22-工程力学作业单（1-2）.docx'
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
    doc = Document(path)
    assert len(doc.tables) == 2 and len(doc.tables[1].rows) == 14
    assert len(doc.inline_shapes) == 1
    assert '2253710052' in '\n'.join(p.text for p in doc.paragraphs)
    with pymupdf.open(path.with_suffix('.preview.pdf')) as pdf:
        assert len(pdf) == 1 and len(pdf[0].get_text()) > 100
    images = list((OUT / '2026-09-22-作业/图').glob('*.png'))
    assert len(images) == 8
    for path in images:
        with Image.open(path) as im:
            im.verify()
    print('PASS：教材双源120dpi像素/文本一致；DOCX可读且14行留白；PDF一页；8张PNG解码正常')


if __name__ == '__main__':
    verify_models()
    verify_artifacts()
