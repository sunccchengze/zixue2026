#!/usr/bin/env python3
"""Offline checks for revised presentation evidence boundaries.

Build first: python 英语大作业/make_deck.py
Run: python -m unittest discover -s scripts -p test_presentation_evidence.py -v
These checks do not replace rendering, listening, or scientific peer review.
"""
import ast
import math
from pathlib import Path
import unittest

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / '英语大作业'


class EvidenceTests(unittest.TestCase):
    def test_source_removes_overclaims(self):
        module = ast.parse((PROJECT / 'make_deck.py').read_text())
        text = '\n'.join(n.value for n in ast.walk(module)
                         if isinstance(n, ast.Constant) and isinstance(n.value, str))
        for claim in ('Hot water kills the virus', 'is the only mechanism',
                      'stops bacteria from growing', 'Verdict: Approved',
                      'the conclusion stands', 'DOABLE IN A KITCHEN'):
            self.assertNotIn(claim, text)
        self.assertIn('BFE: a separate bacterial-count endpoint', text)
        self.assertIn('not medical advice', text)
        self.assertIn('3M 9502', text)
        self.assertIn('3M 1860', text)

    def test_old_roles_are_labelled_historical(self):
        for path in (PROJECT / 'roles').glob('*.md'):
            self.assertIn('历史讲稿（保留溯源）', path.read_text().splitlines()[0])
            self.assertIn('修订讲稿-证据边界版.md', path.read_text().splitlines()[0])

    def test_revised_script_preserves_uncertainty(self):
        text = (PROJECT / '修订讲稿-证据边界版.md').read_text()
        for expected in ('P4负责11–13', '0.075', '0.020', '3M 9502',
                         '3M 1860', 'KF94', 'not approval',
                         'cannot resolve that discrepancy'):
            self.assertIn(expected, text)

    def test_ideal_copper_voltage(self):
        coefficient = 8.314462618 * 298.15 / (2 * 96485.33212)
        self.assertAlmostEqual(coefficient * math.log(10), 0.029580, places=6)
        self.assertAlmostEqual(coefficient * math.log(100), 0.059159, places=6)
        text = (ROOT / '大化大作业/修订讲稿-理论条件版.md').read_text()
        self.assertIn('不是实际实验读数', text)
        self.assertIn('暂撤下定量柱图', text)

    def test_generated_deck(self):
        path = PROJECT / 'build/Mask-Reuse-Qualification-Review.pptx'
        if not path.exists():
            self.skipTest('Run make_deck.py first; build output is intentionally not tracked')
        from pptx import Presentation
        deck = Presentation(path)
        self.assertEqual(len(deck.slides), 13)
        text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes
                         if shape.has_text_frame)
        self.assertNotIn('Hot water kills the virus', text)
        self.assertNotIn('is the only mechanism', text)
        self.assertIn('Limited Evidence, Not Approval', text)
        self.assertIn('孙承泽', text)
        for slide in deck.slides:
            self.assertIn('not medical advice', slide.notes_slide.notes_text_frame.text)


if __name__ == '__main__':
    unittest.main()
