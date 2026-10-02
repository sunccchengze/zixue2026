#!/usr/bin/env python3
"""独立核验9月29日作业：受力模型、Markdown图片渲染、PNG/DOCX/PDF完整性。

仓库根运行：.venv/bin/python scripts/verify_mechanics_homework_20260929.py
依赖同生成器，另需markdown-it-py、python-docx。不调用生成器的自检，
读取绘图用受力模型后独立检查；3-6另用线性方程组重算，不复用解析公式。1-4示意尺寸不是教材题设。
模型核验不能代替逐张人工对照原图，也不声称核对过不存在的第1章官方答案。
"""
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
import random
import subprocess
import sys
from urllib.parse import unquote
import zipfile

from docx import Document
from markdown_it import MarkdownIt
import numpy as np
from PIL import Image
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "工程力学/作业/2026-09-29-作业"
sys.path.insert(0, str(ROOT / "scripts"))
from generate_mechanics_homework_20260929 import models_14, model_36


def check_balance(loads):
    """不使用生成器的平衡检查；对三个不同矩心计算。"""
    dimension = len(loads[0].vector)
    scale_f = max(1, sum(np.linalg.norm(load.vector) for load in loads))
    assert np.linalg.norm(np.sum([load.vector for load in loads], axis=0)) < 1e-10 * scale_f
    origins = [(0, 0), (2.7, -1.9), (-3.2, 4.1)] if dimension == 2 else [(0, 0, 0), (70, -90, 180), (-120, 80, 300)]
    for origin in origins:
        moments = []
        for load in loads:
            r = load.point - origin
            v = load.vector
            moments.append(r[0] * v[1] - r[1] * v[0] if dimension == 2 else np.cross(r, v))
        scale_m = max(1, sum(np.linalg.norm(m) for m in moments))
        assert np.linalg.norm(np.sum(moments, axis=0)) < 1e-10 * scale_m


def pick(body, symbol, node=None):
    matches = [load for load in body.forces if load.symbol == symbol
               and (node is None or np.allclose(load.point, body.nodes[node]))]
    assert len(matches) == 1, (body.title, symbol, node)
    return matches[0]


def opposite(left, right):
    assert np.allclose(left.vector, -right.vector), (left.symbol, right.symbol)


def verify_14():
    expected = {
        "a-ab": ["N_{AB}"] * 2, "a-bc": ["N_{BC}"] * 2,
        "a-pin": ["N_{AB}", "N_{BC}", "P"], "a-whole": ["R_A", "R_C", "P"],
        "b-ab": ["A_x", "A_y", "B_x", "B_y", "T"],
        "b-cd": ["C_x", "C_y", "B_x", "B_y", "D_x", "D_y"],
        "b-h": ["T", "T", "P"], "b-d": ["D_x", "D_y", "T", "T", "T"],
        "b-whole": ["A_x", "A_y", "C_x", "C_y", "P"],
        "c-ab": ["A_x", "A_y", "D_x", "D_y", "N_B", "F_{EB}"],
        "c-ce": ["D_x", "D_y", "F_{EB}", "T", "P"],
        "c-whole": ["A_x", "A_y", "N_B", "T", "P"],
    }
    rng = random.Random(20260929)
    for _ in range(120):
        p = rng.uniform(1, 4000)
        bodies = models_14(p)
        assert set(bodies) == set(expected)
        for key, symbols in expected.items():
            body = bodies[key]
            assert Counter(load.symbol for load in body.forces) == Counter(symbols), key
            check_balance(body.forces)
            for load in body.forces:
                # 铰链的x/y分量必须水平/竖直，不能用两项斜箭头或一项合力替代。
                if load.symbol.endswith("_x"):
                    assert load.vector[1] == 0
                if load.symbol.endswith("_y"):
                    assert load.vector[0] == 0
        for key, symbol in (("a-ab", "N_{AB}"), ("a-bc", "N_{BC}")):
            first, second = bodies[key].forces
            opposite(first, second)
            r = second.point - first.point
            assert abs(r[0] * first.vector[1] - r[1] * first.vector[0]) < 1e-9 * p
            opposite(pick(bodies[key], symbol, "B"), pick(bodies["a-pin"], symbol))
        for left, right, node, symbols in [
            ("b-ab", "b-cd", "B", ("B_x", "B_y")),
            ("b-cd", "b-d", "D", ("D_x", "D_y")),
            ("c-ab", "c-ce", "D", ("D_x", "D_y")),
        ]:
            for symbol in symbols:
                opposite(pick(bodies[left], symbol, node), pick(bodies[right], symbol, node))
        opposite(pick(bodies["c-ab"], "F_{EB}"), pick(bodies["c-ce"], "F_{EB}"))
        for load in bodies["b-d"].forces:
            if load.symbol == "T":
                assert np.isclose(np.linalg.norm(load.vector), p / 2)
        assert np.isclose(np.linalg.norm(pick(bodies["c-ce"], "T").vector), p)
        # 绳力位于轮缘时须沿切线；轴心绳端不是轮缘接触。
        for key, names in (("b-d", ("T",)), ("b-h", ("T",)), ("c-ce", ("T", "P"))):
            center, radius = bodies[key].circles[0]
            for load in bodies[key].forces:
                if load.symbol not in names:
                    continue
                r = load.point - center
                if np.linalg.norm(r) > 1e-10:
                    assert np.isclose(np.linalg.norm(r), radius)
                    assert abs(np.dot(r, load.vector)) < 1e-9 * p
    print("PASS：120组×12个分离体/整体，外力清单、正交分量、绳切线、内部力反向及三个矩心平衡")


def solve_36(p, arm_f, arm_p, yc, yb, yd):
    """独立实现6×5线性系统，不复用生成器的解析公式。"""
    columns = []
    for position, direction in [
        ((arm_f, yc, 0), (0, 0, 1)),
        ((0, 0, 0), (1, 0, 0)), ((0, 0, 0), (0, 0, 1)),
        ((0, yb, 0), (1, 0, 0)), ((0, yb, 0), (0, 0, 1)),
    ]:
        columns.append(np.r_[direction, np.cross(position, direction)])
    matrix = np.column_stack(columns)
    external = np.array((p, 0, 0))
    rhs = -np.r_[external, np.cross((0, yd, arm_p), external)]
    answer, _, rank, _ = np.linalg.lstsq(matrix, rhs, rcond=None)
    assert rank == 5
    assert np.allclose(matrix @ answer, rhs, rtol=1e-9, atol=1e-6)
    return dict(zip(("F", "X_A", "Z_A", "X_B", "Z_B"), answer))


def verify_36():
    rng = random.Random(36)
    samples = [(800., 200., 200., 400., 1000., 1400.)]
    for _ in range(120):
        yb = rng.uniform(700, 1700)
        samples.append((rng.uniform(50, 4000), rng.uniform(100, 400), rng.uniform(100, 400),
                        rng.uniform(.1, .7) * yb, yb, rng.uniform(1.1, 1.8) * yb))
    for parameters in samples:
        loads, given = model_36(*parameters)
        assert Counter(load.symbol for load in loads) == Counter(("F", "P", "X_A", "Z_A", "X_B", "Z_B"))
        independent = solve_36(*parameters)
        for name, value in independent.items():
            assert np.isclose(given[name], value, rtol=1e-9, atol=1e-7), name
        check_balance(loads)
    print("PASS：3-6题设及120组不等偏心距参数，独立6×5系统与三矩心向量平衡均通过")


class Images(HTMLParser):
    def __init__(self):
        super().__init__()
        self.sources = []

    def handle_starttag(self, tag, attrs):
        if tag == "img":
            self.sources.append(dict(attrs)["src"])


def verify_artifacts():
    expected = {
        "参考答案.md": ["图/题1-4.png", "图/解答-1-4-a.png", "图/解答-1-4-b.png",
                    "图/解答-1-4-c.png", "图/题3-6.png", "图/解答-3-6.png"],
        "题面.md": ["图/题1-4.png", "图/题3-6.png"],
    }
    markdown = MarkdownIt("commonmark")
    for filename, sources in expected.items():
        parser = Images()
        parser.feed(markdown.render((OUT / filename).read_text()))
        assert [unquote(src) for src in parser.sources] == sources, (filename, parser.sources)
        for src in parser.sources:
            path = OUT / unquote(src)
            assert path.is_file(), path
            with Image.open(path) as image:
                image.verify()
    # 兼容旧链接的总图同样必须可解码，不能遗留错误图或破损文件。
    for path in (OUT / "图").glob("*.png"):
        with Image.open(path) as image:
            assert image.width >= 300 and image.height >= 200, path
            image.verify()
    print("PASS：答案HTML实际包含6个img、题面包含2个img；全部相对路径存在且PNG可解码")

    # 双源查页，明确区分1-based页码与0-based索引。
    with pymupdf.open(ROOT / "工程力学/框架版教材/工程力学（框架汇编版）.pdf") as book, \
         pymupdf.open(ROOT / "工程力学/资料原件/工程力学（第三版）_1-130.pdf") as original:
        for frame_page, original_page, question in ((31, 34, "1-4"), (78, 81, "3-6")):
            a, b = book[frame_page - 1], original[original_page - 1]
            assert question in a.get_text()
            assert a.get_text() == b.get_text()
            assert a.get_pixmap(dpi=72).samples == b.get_pixmap(dpi=72).samples
        answer = book[288].get_text().split("3-6", 1)[1].split("3-7", 1)[0]
        for value in ("800N", "320N", "480N", "1120N"):
            assert value in answer
    spec_text = (OUT / "作业单.json").read_text()
    assert "PDF31" in spec_text and "PDF78" in spec_text
    print("PASS：1-4/3-6的框架版31/78页与原件34/81页像素/文本相同；答案位于框架版289页")

    path = OUT / "工程力学作业单（1-4、3-6）.docx"
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for value in ("孙承泽", "2253710052", "能动强基2501", "PDF31", "PDF78"):
        assert value in text
    assert len(doc.inline_shapes) == 2
    assert sum(len(table.rows) for table in doc.tables if len(table.columns) == 1) == 14
    audit = subprocess.run([sys.executable, str(ROOT / "scripts/docx_page_audit.py"), str(path)],
                           check=True, capture_output=True, text=True)
    assert "一页放得下" in audit.stdout
    with pymupdf.open(path.with_suffix(".preview.pdf")) as pdf:
        assert len(pdf) == 1
        assert "2253710052" in pdf[0].get_text()
        assert "PDF31" in pdf[0].get_text() and "PDF78" in pdf[0].get_text()
    print("PASS：DOCX可读，2题图/14留白线/署名及新页码正确，XML单页审计与PDF一页检查通过")


if __name__ == "__main__":
    verify_14()
    verify_36()
    verify_artifacts()
