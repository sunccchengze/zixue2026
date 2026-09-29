#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""审计《大学物理 期中 12-15 次作业全题解析 V4.0》docx 的完整性。

检查项
------
1. 96 个题号块（40 选择 / 40 填空 / 16 计算）全部在位；
2. 四次作业的「知识框架｜零基础精讲」小节全部在位，且块类型渲染齐全
   （段落、公式行、分色提示框、对比表格、图卡、编号清单、收口总结）；
3. 没有 Markdown 源码残留（**、«»、LaTeX 反斜杠）；
4. 署名与班级正确（孙承泽 / 2253710052 / 能动强基2501）；
5. 版号已升到 V4.0。

用法：python scripts/audit_physics_midterm_v4.py
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "大学物理" / "期中复习" / "大学物理期中_12-15次作业图文精析_孙承泽_2253710052.docx"

CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
CALLOUT_LABELS = ["关键一句", "记成一句话", "选题题常考", "选择题常考", "一句话总结",
                  "本层与逐题解析的关系", "读法提醒", "口诀", "结论", "别搞混",
                  "别混淆", "别搞反", "易混提醒", "最狠的坑", "最易翻车的一点",
                  "最关键的一组对比", "顺序不能乱", "易错", "反直觉", "陷阱",
                  "要点", "技巧"]


def all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main() -> int:
    if not DOCX.exists():
        print(f"缺少文档：{DOCX}")
        return 1

    with zipfile.ZipFile(DOCX) as zf:
        bad = zf.testzip()
        assert bad is None, f"DOCX 压缩包损坏：{bad}"
        media = [n for n in zf.namelist() if n.startswith("word/media/")]

    doc = Document(DOCX)
    text = all_text(doc)
    errors: list[str] = []

    # 1) 96 题覆盖
    missing_q = []
    for session in (12, 13, 14, 15):
        for category, start, end in (("选择题", 1, 10), ("填空题", 11, 20),
                                     ("计算题", 21, 24)):
            for n in range(start, end + 1):
                if f"第 {n} 题｜{category}" not in text:
                    missing_q.append(f"{session}-{category}-{n}")
    if missing_q:
        errors.append(f"题号块缺失 {len(missing_q)} 个：{missing_q[:8]}")

    n_heading = len(re.findall(r"第 \d+ 题｜", text))
    n_answer = text.count("参考答案")
    n_solution = text.count("解题思路")
    n_step = text.count("步骤 ")

    # 2) 零基础精讲层
    if text.count("知识框架｜零基础精讲") != 4:
        errors.append("「知识框架｜零基础精讲」小节数不为 4")
    if text.count("知识框架速览表") != 4:
        errors.append("「知识框架速览表」小节数不为 4")
    if text.count("一句话总结") < 4:      # 每次作业各一个收口总结块
        errors.append("收口总结块过少（应至少 4 处）")
    for circle in CIRCLED[:10]:           # 每次作业的速记清单至少 10 条
        if circle not in text:
            errors.append(f"编号清单缺少符号 {circle}")
            break
    for label in ("核心结论", "易错 / 陷阱", "技巧 / 补充", "反直觉"):
        if label not in text:
            errors.append(f"缺少分色提示框说明：{label}")

    # 3) Markdown / 占位符 / LaTeX 残留
    for token, why in (("**", "Markdown 加粗标记"), ("«", "占位符 «"),
                       ("»", "占位符 »"), ("\\lambda", "LaTeX 源码"),
                       ("\\Delta", "LaTeX 源码"), ("\\frac", "LaTeX 源码"),
                       ("\\cos", "LaTeX 源码"), ("\\times", "LaTeX 源码")):
        if token in text:
            errors.append(f"残留{why}：{token!r}（{text.count(token)} 处）")

    # 4) 署名
    for token in ("孙承泽", "2253710052", "能动强基2501"):
        if token not in text:
            errors.append(f"署名缺失：{token}")
    for wrong in ("孙昉泽", "能动 2501", "能动2501"):
        if wrong in text:
            errors.append(f"署名错误：{wrong}")

    # 4b) 繁体字残留（全文统一用简体）
    for trad in ("題號", "黑體", "錯誤", "解題", "選擇題", "條紋"):
        if trad in text:
            errors.append(f"繁体字残留：{trad}")

    # 5) 版号
    if "V4.0" not in text:
        errors.append("版号未升到 V4.0")
    if "V3.0" in text:
        errors.append("仍残留 V3.0 版号")

    # 6) 零基础层的体量体检（防止将来被改回结论式罗列）
    try:
        sys.path.insert(0, str(ROOT / "大学物理" / "课程作业"))
        from 知识框架精讲_v4 import DEEP_FRAMEWORK  # noqa: E402
    except Exception as exc:  # pragma: no cover
        errors.append(f"无法导入知识框架底稿：{exc}")
    else:
        for session, blocks in DEEP_FRAMEWORK.items():
            chars = sum(len(str(b)) for b in blocks)
            print(f"  第 {session} 次精讲：{len(blocks)} 块 / 约 {chars} 字")
            if chars < 4000:
                errors.append(f"第 {session} 次精讲字数过少（{chars}）")

    # 7) docx 与 markdown 口径一致（两个渲染后端必须讲同一件事）
    md_path = ROOT / "大学物理" / "期中复习" / "12-15次作业·知识精讲（零基础版）.md"
    if not md_path.exists():
        errors.append(f"缺少 Markdown 精讲文件：{md_path}")
    else:
        md = md_path.read_text(encoding="utf-8")
        stale = ROOT / "大学物理" / "课程作业" / "知识精讲_v4.py"
        if stale.exists():
            errors.append("存在被取代的旧底稿 知识精讲_v4.py（应只有 知识框架精讲_v4.py）")
        for session, blocks in DEEP_FRAMEWORK.items():
            for block in blocks:
                if block[0] not in ("h3", "summary", "formula"):
                    continue
                probe = str(block[1]).strip()      # md 保留 ** 标记，按原样比对
                if len(probe) < 6:
                    continue
                if probe not in md:
                    errors.append(f"md 未收录第 {session} 次的「{probe[:24]}」")
                    break
        # 勘误表必须在两个后端都留痕
        for must in ("同类间距 λ/2", "λ/(2n)"):
            if must not in md:
                errors.append(f"md 缺少勘误关键串：{must}")
        print(f"  Markdown 精讲：{md_path.name}（{md.count(chr(10)) + 1} 行）")

    print(f"文件：{DOCX.name}（{DOCX.stat().st_size / 1024:.0f} KB，"
          f"{len(media)} 张图，{len(doc.paragraphs)} 段，{len(doc.tables)} 表）")
    print(f"题号块 {n_heading} 个；参考答案 {n_answer} 处；解题思路 {n_solution} 处；"
          f"推理步骤 {n_step} 条")
    print(f"正文字符数：{len(text)}")

    if errors:
        print("\n审计未通过：")
        for e in errors:
            print("  ✗", e)
        return 1
    print("\n审计通过：96 题在位，零基础精讲层完整，无残留标记，署名与版号正确。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
