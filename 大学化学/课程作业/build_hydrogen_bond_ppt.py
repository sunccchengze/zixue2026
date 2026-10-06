from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pathlib import Path

OUT = Path(__file__).with_name('课程作业-氢键如何塑造水与生命.pptx')
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# Palette: deep ink, warm paper, turquoise, coral.
BG = 'F5F3EE'; INK = '172B3A'; MUTED = '58707E'; TEAL = '1F9D91'; PALE = 'DDF2EE'; CORAL = 'F07962'; GOLD = 'EABF62'; WHITE = 'FFFFFF'; LINE = 'D7E0DF'; NAVY = '102532'; SOFT = 'EAF0EF'
FONT = 'Microsoft YaHei'

def rgb(h): return RGBColor.from_string(h)
def set_bg(slide, color=BG):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(color)
def shape(slide, kind, x,y,w,h, fill=None, line=None, radius=False):
    s=slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb=rgb(line); s.line.width=Pt(1.2)
    else: s.line.fill.background()
    return s
def text(slide, value, x,y,w,h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, font=FONT):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=value; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    r._r.get_or_add_rPr().set(qn('a:ea'), font)
    return box
def rich_lines(slide, items, x,y,w,h, size=16, color=INK, gap=8, bullet=False):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.margin_top=Inches(.05); tf.margin_bottom=Inches(.02)
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap); p.level=0
        lead, body = item if isinstance(item,tuple) else ('',item)
        if bullet: lead='• '+lead if lead else '• '
        if lead:
            r=p.add_run(); r.text=lead; r.font.name=FONT; r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=rgb(TEAL); r._r.get_or_add_rPr().set(qn('a:ea'),FONT)
        r=p.add_run(); r.text=body; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=rgb(color); r._r.get_or_add_rPr().set(qn('a:ea'),FONT)
    return box
def line(slide,x1,y1,x2,y2,color=LINE,width=1.5,dash=None):
    ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); ln.line.color.rgb=rgb(color); ln.line.width=Pt(width)
    if dash: ln.line.dash_style=dash
    return ln
def circle(slide,cx,cy,d,label,fill,color=WHITE,fs=19):
    s=shape(slide,MSO_SHAPE.OVAL,cx-d/2,cy-d/2,d,d,fill)
    text(slide,label,cx-d/2,cy-d/2+.01,d,d-.02,fs,color,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,0)
    return s
def footer(slide,n,dark=False):
    c='A9BBC2' if dark else MUTED
    text(slide,'大学化学 · 课堂分享作业',.55,7.12,4,.2,9,c)
    text(slide,f'{n:02d} / 09',11.95,7.10,.8,.22,9,c,False,PP_ALIGN.RIGHT)
def title(slide,eyebrow,headline,sub=None):
    text(slide,eyebrow.upper(),.72,.42,11,.28,10,TEAL,True)
    text(slide,headline,.7,.83,12, .62,29,INK,True)
    if sub: text(slide,sub,.72,1.52,11.8,.48,14,MUTED)
def card(slide,x,y,w,h,fill=WHITE):
    s=shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,fill)
    s.adjustments[0]=.12
    return s

# 1 Cover
s=prs.slides.add_slide(BLANK); set_bg(s,NAVY)
text(s,'大学化学 · 概念分享',.8,.58,5,.32,12,'8ED4C9',True)
text(s,'氢键',.78,1.25,5,.9,46,WHITE,True)
text(s,'一条看不见的相互作用，\n怎样塑造水与生命？',.82,2.25,6.4,1.42,27,'E9F1F1',True)
text(s,'从 D—H···A 的微观图景，到水的反常性质、DNA 与蛋白质',.85,4.08,6.5,.62,15,'B8C9CF')
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.82,5.14,5.6,.78,'203B49')
text(s,'孙承泽  |  能动强基2501  |  2253710052',1.02,5.34,5.2,.3,13,WHITE,True)
# molecule motif right
circle(s,9.45,2.45,.78,'O',TEAL,WHITE,22); circle(s,8.63,3.15,.43,'H',WHITE,INK,13); circle(s,10.27,3.15,.43,'H',WHITE,INK,13)
line(s,9.15,2.73,8.82,2.97,'B9D6D3',2); line(s,9.75,2.73,10.08,2.97,'B9D6D3',2)
circle(s,11.75,4.38,.78,'O',CORAL,WHITE,22); circle(s,10.93,5.08,.43,'H',WHITE,INK,13); circle(s,12.57,5.08,.43,'H',WHITE,INK,13)
line(s,11.45,4.66,11.12,4.9,'B9D6D3',2); line(s,12.05,4.66,12.38,4.9,'B9D6D3',2)
line(s,8.84,3.34,10.65,4.67,GOLD,2,MSO_LINE_DASH_STYLE.DASH)
text(s,'氢键',9.35,3.88,1.1,.35,15,GOLD,True,PP_ALIGN.CENTER)
text(s,'D—H···A',9.26,6.13,2.3,.45,21,WHITE,True,PP_ALIGN.CENTER)
footer(s,1,True)

# 2 Hook
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'01 · 先看反常','水为什么“不按常理出牌”？','两个日常事实，把问题指向分子间作用力。')
card(s,.75,2.22,5.7,3.62,WHITE); card(s,6.85,2.22,5.7,3.62,WHITE)
text(s,'01  沸点偏高',1.08,2.56,4.9,.42,20,TEAL,True)
text(s,'同族氢化物里，水的沸点显著高于 H₂S。',1.08,3.16,4.85,.88,22,INK,True)
text(s,'只看分子量与色散力，解释不了全部差异。\n水分子之间还有方向性较强的吸引。',1.08,4.34,4.8,.9,15,MUTED)
text(s,'02  冰会浮在水上',7.18,2.56,4.9,.42,20,CORAL,True)
text(s,'结冰后，水的密度反而变小。',7.18,3.16,4.85,.88,22,INK,True)
text(s,'冰的氢键网络形成较开放的排列；融化时，\n部分网络塌缩，分子能排得更紧。',7.18,4.34,4.85,.9,15,MUTED)
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,2.25,6.13,8.85,.55,PALE)
text(s,'主问题：一种分子间吸引，如何从“微观很小”累积成“宏观很大”？',2.45,6.24,8.45,.27,15,INK,True,PP_ALIGN.CENTER)
footer(s,2)

# 3 Ingredients
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'02 · 看分子','氢键要先凑齐“供体”与“受体”','典型教学模型：D—H···A；重点不是看到 H 就判氢键。')
card(s,.8,2.25,5.65,3.7,WHITE); card(s,6.85,2.25,5.65,3.7,WHITE)
text(s,'供体 donor：D—H',1.16,2.61,4.9,.4,20,TEAL,True)
text(s,'D 与 H 之间的共价键极化，让 H 带有较强的部分正电。',1.16,3.18,4.85,.86,17,INK)
# D-H polarity sketch
circle(s,2.45,4.75,.66,'O',TEAL,WHITE,20); circle(s,3.72,4.75,.45,'H',WHITE,INK,16)
line(s,2.77,4.75,3.49,4.75,INK,3); text(s,'δ−',2.23,5.15,.5,.3,12,TEAL,True); text(s,'δ+',3.5,5.15,.5,.3,12,CORAL,True)
text(s,'受体 acceptor：富电子区域',7.22,2.61,4.9,.4,20,CORAL,True)
text(s,'常见为带孤对电子的原子；也可能是 π 电子等富电子区域。',7.22,3.18,4.85,.9,17,INK)
circle(s,9.1,4.75,.66,'O',CORAL,WHITE,20); text(s,'••',9.42,4.45,.45,.3,20,INK,True); text(s,'孤对电子',9.98,4.62,1.6,.3,13,MUTED)
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,2.0,6.22,9.35,.56,SOFT)
text(s,'常见强氢键组合是 O—H、N—H、F—H 作供体；受体范围不只限于 O / N / F。',2.18,6.34,9.0,.28,13,INK,True,PP_ALIGN.CENTER)
footer(s,3)

# 4 Definition
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'03 · 符号落地','把氢键写成：D—H···A','实线是共价键；点线表示氢键相互作用，不是另一根普通共价键。')
card(s,.9,2.2,11.55,3.95,WHITE)
# large schematic
circle(s,3.15,4.08,.92,'D',TEAL,WHITE,24); circle(s,5.25,4.08,.7,'H',GOLD,INK,21); circle(s,8.3,4.08,.92,'A',CORAL,WHITE,24)
line(s,3.61,4.08,4.90,4.08,INK,3); line(s,5.60,4.08,7.82,4.08,TEAL,2,MSO_LINE_DASH_STYLE.DASH)
text(s,'供体部分',2.1,2.73,3.8,.3,15,TEAL,True,PP_ALIGN.CENTER)
text(s,'受体（富电子区）',7.0,2.73,2.7,.3,15,CORAL,True,PP_ALIGN.CENTER)
text(s,'D—H 共价键',3.3,4.55,2.0,.35,13,MUTED,False,PP_ALIGN.CENTER)
text(s,'H···A 氢键',5.7,4.55,2.1,.35,13,TEAL,True,PP_ALIGN.CENTER)
text(s,'相互作用常有方向性：D—H···A 越接近线性，通常越有利；但不能把某个角度阈值当成所有体系的硬性定义。',1.35,5.35,10.5,.58,16,INK,False,PP_ALIGN.CENTER)
text(s,'例：水分子之间 O—H···O；DNA 碱基之间 N—H···O / N。',1.55,6.48,10.0,.34,14,MUTED,False,PP_ALIGN.CENTER)
footer(s,4)

# 5 Physical picture
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'04 · 追问本质','它不是“一种神秘胶水”','氢键是多种物理贡献共同形成的吸引；教材入门可先抓住静电图像。')
# left stacked model
card(s,.78,2.2,7.35,4.32,WHITE)
text(s,'一个够用、但不冒充完整的图像',1.12,2.55,6.5,.4,19,INK,True)
rich_lines(s,[('静电吸引：','极化的 Hδ+ 与受体富电子区域相互吸引，是直觉起点。'),('电子重排：','受体电子云可被极化；部分体系还有供体—受体电荷转移贡献。'),('方向性：','相对取向改变轨道重叠与静电分布，因此氢键不是各向同性的“球形粘住”。')],1.12,3.18,6.55,2.75,15,gap=13)
# right contrast
card(s,8.48,2.2,4.08,4.32,NAVY)
text(s,'避免两个误区',8.82,2.57,3.3,.4,19,'8ED4C9',True)
text(s,'×  把氢键说成普通共价键\n\n×  认为只有静电力、没有其他贡献\n\n×  认为任意 X—H 都能形成强氢键',8.82,3.23,3.25,2.2,15,WHITE)
text(s,'定义看证据；模型看尺度。',8.82,5.83,3.3,.35,14,GOLD,True)
text(s,'能量不是固定常数：受分子、几何、溶剂与环境影响，存在连续谱。',1.05,6.7,11.2,.26,12,MUTED,False,PP_ALIGN.CENTER)
footer(s,5)

# 6 Macro
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'05 · 从微观到宏观','单个不必惊人，网络可以改变世界','液态水中氢键持续形成、断裂与重组；宏观性质来自动态网络，不是静态“每分子固定绑四根”。')
card(s,.8,2.18,5.7,3.95,WHITE); card(s,6.82,2.18,5.7,3.95,WHITE)
text(s,'液态水：动态网络',1.15,2.52,4.9,.38,20,TEAL,True)
# connected network
pts=[(2.0,3.72),(3.28,3.72),(4.55,3.72),(2.64,4.82),(3.92,4.82),(5.2,4.82)]
for a,b in [(0,1),(1,2),(0,3),(1,3),(1,4),(2,4),(3,4),(4,5)]: line(s,*pts[a],*pts[b],GOLD,1.5,MSO_LINE_DASH_STYLE.DASH)
for i,(x,y) in enumerate(pts): circle(s,x,y,.45,'O',TEAL,WHITE,13)
text(s,'重排与断裂需要能量 → 沸点、热容等性质被显著改变。',1.12,5.37,5.0,.47,14,INK,False,PP_ALIGN.CENTER)
text(s,'冰：较开放的网络',7.16,2.52,4.9,.38,20,CORAL,True)
# hexagon network schematic
import math
cx,cy=9.68,4.25; R=1.08
hexpts=[]
for i in range(6):
    ang=math.radians(-90+i*60); hexpts.append((cx+R*math.cos(ang),cy+R*math.sin(ang)))
for i in range(6): line(s,*hexpts[i],*hexpts[(i+1)%6],GOLD,1.7,MSO_LINE_DASH_STYLE.DASH)
for x,y in hexpts: circle(s,x,y,.38,'O',CORAL,WHITE,11)
text(s,'开放晶格占空间更多 → 同质量体积更大 → 冰密度更小。',7.1,5.37,5.0,.47,14,INK,False,PP_ALIGN.CENTER)
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,2.07,6.38,9.2,.48,PALE)
text(s,'宏观效应的关键：相互作用 × 数量 × 几何网络 × 热运动。',2.25,6.48,8.85,.25,14,INK,True,PP_ALIGN.CENTER)
footer(s,6)

# 7 Biology
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'06 · 走进生命','DNA 与蛋白质：氢键帮忙“认准形状”','氢键提供方向性与可逆性；但它不是生命大分子稳定性的唯一来源。')
card(s,.8,2.2,5.72,3.98,WHITE); card(s,6.8,2.2,5.72,3.98,WHITE)
text(s,'DNA 碱基配对',1.14,2.55,4.9,.4,20,TEAL,True)
# simplified pairs
text(s,'A  ···  T',1.35,3.35,4.55,.5,27,INK,True,PP_ALIGN.CENTER)
text(s,'G  ···  C',1.35,4.18,4.55,.5,27,INK,True,PP_ALIGN.CENTER)
text(s,'典型 Watson–Crick 配对：A—T 两条氢键；G—C 三条。\n互补供受体图样参与分子识别。',1.25,5.08,4.85,.72,14,MUTED,False,PP_ALIGN.CENTER)
text(s,'蛋白质折叠',7.14,2.55,4.9,.4,20,CORAL,True)
# stylized helix / folded backbone
for i in range(6):
    y=3.43+i*.33; x=7.65+(i%2)*.42
    circle(s,x,y,.15,'',TEAL,WHITE,4)
    circle(s,x+1.25,y,.15,'',CORAL,WHITE,4)
    line(s,x+.09,y,x+1.16,y,GOLD,1.5,MSO_LINE_DASH_STYLE.DASH)
text(s,'骨架 C=O 与 N—H 可形成氢键，帮助稳定 α 螺旋、β 折叠等二级结构。',7.12,5.48,4.9,.52,14,MUTED,False,PP_ALIGN.CENTER)
text(s,'提醒：DNA 还受碱基堆积等作用稳定；蛋白质还涉及疏水效应、静电与范德华作用。',1.0,6.55,11.3,.3,13,INK,True,PP_ALIGN.CENTER)
footer(s,7)

# 8 Check
s=prs.slides.add_slide(BLANK); set_bg(s,NAVY)
text(s,'07 · 自测闭环',.78,.48,4,.3,11,'8ED4C9',True)
text(s,'不看上一页，试着解释',.76,.95,11.5,.65,31,WHITE,True)
qs=[('01','乙醇与二甲醚分子量相同，乙醇沸点更高；氢键如何解释？'),('02','水分子里哪个原子可作氢键供体？哪个可作受体？'),('03','“氢键越多，物质沸点一定越高”这句话完整吗？还要考虑什么？')]
for i,(num,q) in enumerate(qs):
    y=2.0+i*1.28
    shape(s,MSO_SHAPE.OVAL,.88,y,.52,.52,TEAL)
    text(s,num,.88,y+.1,.52,.25,12,WHITE,True,PP_ALIGN.CENTER)
    text(s,q,1.7,y-.02,10.6,.72,19,WHITE,True)
text(s,'思考提示：供体 / 受体 · 分子间作用力 · 分子结构与形状 · 温度下的动态网络',1.72,6.16,10.6,.55,13,'B8C9CF')
footer(s,8,True)

# 9 takeaway
s=prs.slides.add_slide(BLANK); set_bg(s); title(s,'08 · 一句话收束','氢键：局部定向吸引，集体塑造宏观结构','记住这条因果链，就能从符号走回真实物质。')
steps=[('极化','D—H 让 H 带部分正电'),('识别','H···A 富电子区相互吸引'),('成网','方向性与环境决定排列'),('显性','沸点、密度、生物结构改变')]
xs=[.9,4.0,7.1,10.2]
for i,(head,body) in enumerate(steps):
    card(s,xs[i],2.7,2.25,1.75,WHITE)
    text(s,head,xs[i]+.14,2.95,1.95,.38,19,TEAL if i<3 else CORAL,True,PP_ALIGN.CENTER)
    text(s,body,xs[i]+.14,3.52,1.95,.62,13,INK,False,PP_ALIGN.CENTER)
    if i<3:
        text(s,'→',xs[i]+2.38,3.25,.6,.4,24,GOLD,True,PP_ALIGN.CENTER)
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,1.05,5.05,11.15,.82,NAVY)
text(s,'下一步：用氢键与偶极—偶极、色散力对照，解释“结构相似，性质为何不同”。',1.35,5.28,10.55,.34,16,WHITE,True,PP_ALIGN.CENTER)
text(s,'参考：Arunan et al., Pure Appl. Chem. 83 (2011), 1637–1641, IUPAC Recommendations. DOI: 10.1351/PAC-REC-10-01-02',.88,6.42,11.7,.35,9,MUTED,False,PP_ALIGN.CENTER)
footer(s,9)

prs.core_properties.title = '氢键：一条看不见的相互作用，怎样塑造水与生命？'
prs.core_properties.subject = '大学化学课程分享作业'
prs.core_properties.author = '孙承泽'
prs.core_properties.keywords = '氢键, D-H···A, 水, DNA, 蛋白质'
prs.save(OUT)
print(OUT)
